#!/usr/bin/env python3
"""Deterministic structural checks plus advisory quality checks for editable PPTX decks."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE


PLACEHOLDER = re.compile(
    r"(?:\[(?:name|website|contact|company|title)\]|\b(?:TBD|TODO)\b|"
    r"待补充|待填写|请补充|请填写|此处(?:补充|填写)|占位文本|示例文本)",
    re.IGNORECASE,
)
SERIALIZED_CONFIG_PAYLOAD = re.compile(
    r"(?:^|\s)[\[{]\s*['\"]?(?:name|owner)['\"]?\s*:\s*"
    r"[^{}\[\]\r\n]{0,180}['\"]?(?:section|tagline|action)['\"]?\s*:",
    re.IGNORECASE,
)
LATIN_MOJIBAKE = re.compile(r"(?:\ufffd|Ã.|Â.|â(?:€|€™|€“|€”|€¦)|ðŸ)")
CJK_MOJIBAKE_TOKENS = (
    "锛", "銆", "鈥", "鈫", "鍒", "瀹", "鍔", "绔", "璇", "鐨", "鍙", "浠",
    "鏂", "缁", "鎴", "娴", "绠", "鐢", "妫", "瑙", "纭", "杈", "寮", "鎵", "璐",
)
INTEGER = re.compile(r"^-?\d+$")
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
EMU_PER_INCH = 914_400


def weighted_text_units(text: str) -> float:
    """Approximate line width across CJK and Latin text without host fonts."""

    units = 0.0
    for char in text:
        if char == "\n":
            continue
        if char.isspace():
            units += 0.3
        elif ord(char) > 127:
            units += 1.0
        elif char.isalnum():
            units += 0.55
        else:
            units += 0.4
    return units


def likely_text_overflow(text: str, width: int, height: int, font_size: float) -> bool:
    """Conservatively flag boxes whose text cannot fit at the authored size."""

    if len(text) < 20 or width <= 0 or height <= 0 or font_size <= 0:
        return False
    width_pt = width / EMU_PER_INCH * 72.0
    height_pt = height / EMU_PER_INCH * 72.0
    chars_per_line = max(1.0, width_pt / (font_size * 0.92))
    lines = max(1, math.floor(height_pt / (font_size * 1.15)))
    explicit_lines = max(1, text.count("\n") + 1)
    if explicit_lines > lines:
        return True
    capacity = chars_per_line * lines * 1.15
    return weighted_text_units(text) > capacity


def likely_mojibake(text: str) -> bool:
    if LATIN_MOJIBAKE.search(text):
        return True
    return sum(text.count(token) for token in CJK_MOJIBAKE_TOKENS) >= 3


def shape_kind(shape) -> str:
    """Return a content-agnostic shape class for layout comparison."""

    if bool(getattr(shape, "has_chart", False)):
        return "chart"
    if bool(getattr(shape, "has_table", False)):
        return "table"
    shape_type = getattr(shape, "shape_type", None)
    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape_type == MSO_SHAPE_TYPE.LINE:
        return "line"
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(shape, "has_text_frame", False) and getattr(shape, "text", "").strip():
        return "text"
    if shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        try:
            auto_shape_type = shape.auto_shape_type
        except (KeyError, ValueError):
            auto_shape_type = "unknown"
        return f"auto:{auto_shape_type}"
    return "shape"


def safe_auto_shape_type(shape):
    try:
        return shape.auto_shape_type
    except (AttributeError, KeyError, ValueError):
        return None


def layout_signature(slide, width: int, height: int) -> Counter[tuple[object, ...]]:
    """Describe body geometry while ignoring text, colors, and repeated chrome.

    The coarse normalized grid tolerates small implementation drift but still
    catches a reference deck being reused as a slide-by-slide layout map.
    """

    signature: Counter[tuple[object, ...]] = Counter()
    for shape in slide.shapes:
        left = int(getattr(shape, "left", 0) or 0)
        top = int(getattr(shape, "top", 0) or 0)
        shape_width = int(getattr(shape, "width", 0) or 0)
        shape_height = int(getattr(shape, "height", 0) or 0)
        area = max(0, shape_width) * max(0, shape_height)
        if area >= width * height * 0.92:
            continue

        x = left / max(1, width)
        y = top / max(1, height)
        w = shape_width / max(1, width)
        h = shape_height / max(1, height)

        # Page numbers, footers, and compact title rails establish a visual
        # system but must not make otherwise unrelated pages look identical.
        if y >= 0.89:
            continue
        if y <= 0.13 and h <= 0.12:
            continue
        if w < 0.012 and h < 0.012:
            continue

        token = (
            shape_kind(shape),
            round(x * 24),
            round(y * 14),
            round(w * 24),
            round(h * 14),
        )
        signature[token] += 1
    return signature


def layout_similarity(
    left: Counter[tuple[object, ...]], right: Counter[tuple[object, ...]]
) -> float:
    intersection = sum((left & right).values())
    return intersection / max(1, sum(left.values()), sum(right.values()))


def layout_fingerprint(signature: Counter[tuple[object, ...]]) -> str:
    payload = "|".join(
        f"{token}:{count}" for token, count in sorted(signature.items(), key=lambda item: repr(item[0]))
    )
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()[:12]


def deck_layout_signatures(path: Path) -> list[Counter[tuple[object, ...]]]:
    deck = Presentation(str(path))
    width, height = int(deck.slide_width), int(deck.slide_height)
    return [layout_signature(slide, width, height) for slide in deck.slides]


def layout_originality_report(
    signatures: list[Counter[tuple[object, ...]]],
    *,
    style_reference: Path | None,
    allow_repeated_layouts: bool,
) -> tuple[dict[str, object], list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    interior = [
        index
        for index in range(1, max(1, len(signatures) - 1))
        if sum(signatures[index].values()) >= 6
    ]

    repeated_pairs: list[dict[str, object]] = []
    for offset, left_index in enumerate(interior):
        for right_index in interior[offset + 1 :]:
            similarity = layout_similarity(signatures[left_index], signatures[right_index])
            if similarity >= 0.94:
                repeated_pairs.append(
                    {
                        "slide_a": left_index + 1,
                        "slide_b": right_index + 1,
                        "similarity": round(similarity, 3),
                        "adjacent": right_index == left_index + 1,
                    }
                )

    adjacent_repeats = [pair for pair in repeated_pairs if pair["adjacent"]]
    repeated_slides = {
        int(pair[key])
        for pair in repeated_pairs
        for key in ("slide_a", "slide_b")
    }
    repeated_message = ""
    if adjacent_repeats:
        pair_text = ", ".join(
            f"{pair['slide_a']}-{pair['slide_b']} ({pair['similarity']:.2f})"
            for pair in adjacent_repeats[:6]
        )
        repeated_message = (
            "adjacent content slides reuse the same body geometry: "
            f"{pair_text}; redesign the focal anchor or composition instead of changing only copy/colors"
        )
    elif len(repeated_slides) >= 3:
        repeated_message = (
            "one body-layout skeleton is reused across multiple content slides: "
            f"{', '.join(map(str, sorted(repeated_slides)))}; use page-specific compositions"
        )
    if repeated_message:
        if allow_repeated_layouts:
            warnings.append(repeated_message)
        else:
            errors.append(repeated_message)

    reference_report: dict[str, object] | None = None
    if style_reference is not None:
        if not style_reference.is_file():
            errors.append(f"style reference not found: {style_reference}")
        else:
            reference_signatures = deck_layout_signatures(style_reference)
            reference_interior = [
                index
                for index in range(1, max(1, len(reference_signatures) - 1))
                if sum(reference_signatures[index].values()) >= 6
            ]
            matches: list[dict[str, object]] = []
            for candidate_index in interior:
                if not reference_interior:
                    break
                scored = [
                    (
                        layout_similarity(
                            signatures[candidate_index], reference_signatures[reference_index]
                        ),
                        reference_index,
                    )
                    for reference_index in reference_interior
                ]
                similarity, reference_index = max(scored)
                if similarity >= 0.82:
                    matches.append(
                        {
                            "slide": candidate_index + 1,
                            "reference_slide": reference_index + 1,
                            "similarity": round(similarity, 3),
                            "same_ordinal": candidate_index == reference_index,
                        }
                    )

            eligible_count = len(interior)
            clone_limit = max(3, math.ceil(eligible_count * 0.35))
            reference_report = {
                "path": str(style_reference),
                "eligible_content_slides": eligible_count,
                "clone_limit": clone_limit,
                "matched_slides": matches,
                "match_ratio": round(len(matches) / max(1, eligible_count), 3),
                "same_ordinal_matches": sum(bool(item["same_ordinal"]) for item in matches),
            }
            if len(matches) >= clone_limit:
                match_text = ", ".join(
                    f"{item['slide']}→{item['reference_slide']} ({item['similarity']:.2f})"
                    for item in matches[:8]
                )
                errors.append(
                    "style-reference layout cloning detected on "
                    f"{len(matches)} of {eligible_count} content slides: {match_text}; "
                    "retain design DNA (palette/type/grid) but rebuild slide architecture from the new content"
                )

    return (
        {
            "policy": "recurring" if allow_repeated_layouts else "audience-facing",
            "internal_repeated_pairs": repeated_pairs,
            "style_reference": reference_report,
        },
        errors,
        warnings,
    )


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def resolve_relationship(owner_rels: str, target: str) -> str:
    rel_path = PurePosixPath(owner_rels)
    owner_dir = rel_path.parent.parent
    candidate = PurePosixPath(target.lstrip("/")) if target.startswith("/") else owner_dir / target
    parts: list[str] = []
    for part in candidate.parts:
        if part == "..":
            if parts:
                parts.pop()
        elif part not in ("", "."):
            parts.append(part)
    return "/".join(parts)


def package_errors(path: Path) -> list[str]:
    errors: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
            errors.append("missing required OOXML presentation parts")

        for name in sorted(names):
            if not name.endswith((".xml", ".rels")):
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError as exc:
                errors.append(f"malformed XML in {name}: {exc}")
                continue

            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                ids = [
                    node.attrib.get("id", "")
                    for node in root.iter()
                    if local_name(node.tag) == "cNvPr"
                ]
                duplicate_ids = [key for key, count in Counter(ids).items() if key and count > 1]
                if duplicate_ids:
                    errors.append(f"duplicate shape ids in {name}: {', '.join(duplicate_ids)}")

                for node in root.iter():
                    for attr in ("x", "y", "cx", "cy"):
                        value = node.attrib.get(attr)
                        if value is not None and not INTEGER.fullmatch(value):
                            errors.append(
                                f"fractional/invalid EMU in {name}: {local_name(node.tag)}.{attr}={value}"
                            )

            if name.endswith(".rels"):
                for relation in root.findall(f"{REL_NS}Relationship"):
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = resolve_relationship(name, target)
                    if resolved and resolved not in names:
                        errors.append(
                            f"missing relationship target in {name}: {relation.attrib.get('Id', '?')} -> {resolved}"
                        )
    return errors


def presentation_report(
    path: Path,
    *,
    style_reference: Path | None = None,
    allow_repeated_layouts: bool = False,
    min_body_font_size: float = 0.0,
    min_title_font_size: float = 0.0,
) -> tuple[list[dict[str, object]], list[str], list[str], dict[str, object]]:
    errors: list[str] = []
    warnings: list[str] = []
    slides: list[dict[str, object]] = []
    signatures: list[Counter[tuple[object, ...]]] = []
    deck = Presentation(str(path))
    width, height = int(deck.slide_width), int(deck.slide_height)
    card_heavy_slides: list[int] = []
    table_slides: list[int] = []

    if not deck.slides:
        errors.append("presentation has no slides")

    for number, slide in enumerate(deck.slides, 1):
        signature = layout_signature(slide, width, height)
        signatures.append(signature)
        texts: list[str] = []
        visual_shapes = 0
        out_of_bounds: list[str] = []
        tiny_text_boxes: list[str] = []
        overflow_text_boxes: list[str] = []
        text_geometries: list[tuple[int, int, int, int, str]] = []
        overlapping_text_pairs: list[tuple[str, str]] = []
        rounded_rectangles = 0
        charts = 0
        tables = 0
        pictures = 0
        serialized_payloads: list[str] = []
        mojibake_payloads: list[str] = []
        below_requested_body_floor: list[str] = []
        title_region_sizes: list[float] = []
        all_explicit_sizes: list[float] = []
        body_text_chars = 0
        shapes = list(slide.shapes)
        for shape_index, shape in enumerate(shapes):
            left = int(getattr(shape, "left", 0) or 0)
            top = int(getattr(shape, "top", 0) or 0)
            shape_width = int(getattr(shape, "width", 0) or 0)
            shape_height = int(getattr(shape, "height", 0) or 0)
            if left < 0 or top < 0 or left + shape_width > width or top + shape_height > height:
                out_of_bounds.append(getattr(shape, "name", "unnamed shape"))

            if getattr(shape, "has_text_frame", False):
                text = " ".join(getattr(shape, "text", "").split())
                if text:
                    texts.append(text)
                    compact = " ".join(text.split())
                    dict_payload = (
                        compact.startswith(("[{'text':", '[{"text":'))
                        and ("'size':" in compact or '"size":' in compact)
                        and ("'color':" in compact or '"color":' in compact)
                        and ("'bold':" in compact or '"bold":' in compact)
                    )
                    enum_payload = (
                        "<PP_PARAGRAPH_ALIGNMENT." in compact
                        or "<MSO_VERTICAL_ANCHOR." in compact
                    )
                    config_payload = bool(SERIALIZED_CONFIG_PAYLOAD.search(compact))
                    if dict_payload or enum_payload or config_payload:
                        serialized_payloads.append(text[:80])
                    if likely_mojibake(text):
                        mojibake_payloads.append(text[:80])
                    if PLACEHOLDER.search(text):
                        warnings.append(f"placeholder on slide {number}: {text[:80]}")
                    explicit_sizes: list[float] = []
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.font.size is not None:
                            explicit_sizes.append(float(paragraph.font.size.pt))
                        explicit_sizes.extend(
                            float(run.font.size.pt)
                            for run in paragraph.runs
                            if run.font.size is not None
                        )
                    all_explicit_sizes.extend(explicit_sizes)
                    if top < int(height * 0.32) and shape_width >= int(width * 0.30):
                        title_region_sizes.extend(explicit_sizes)
                    # A few small eyebrow/footer labels are legitimate. A page
                    # made from many <=12 pt text boxes is not presentation-
                    # readable and was the main failure mode of one-shot
                    # python-pptx decks. Keep this structural gate deliberately
                    # conservative so ordinary 16-18 pt body copy passes.
                    is_footer = top >= int(height * 0.90)
                    is_header_chrome = (
                        top < int(height * 0.10)
                        and shape_height <= int(height * 0.12)
                        and explicit_sizes
                        and max(explicit_sizes) <= 12.0
                    )
                    is_decorative_index = (
                        bool(re.fullmatch(r"(?:0?\d{1,2}|[IVX]{1,5})", text, re.IGNORECASE))
                        and shape_height <= int(height * 0.08)
                        and explicit_sizes
                        and max(explicit_sizes) <= 12.0
                    )
                    is_title_region = (
                        top < int(height * 0.22)
                        and shape_width >= int(width * 0.30)
                    )
                    is_compact_annotation = (
                        len(text) <= 16
                        and shape_height <= int(height * 0.10)
                        and explicit_sizes
                        and max(explicit_sizes) >= 10.0
                    )
                    if (
                        len(text) >= 8
                        and not is_footer
                        and not is_header_chrome
                        and not is_decorative_index
                        and shape_width > 0
                        and shape_height > 0
                    ):
                        text_geometries.append(
                            (left, top, shape_width, shape_height, text[:48])
                        )
                    if (
                        not is_footer
                        and not is_header_chrome
                        and not is_title_region
                        and not is_decorative_index
                    ):
                        body_text_chars += len(text)
                    if (
                        min_body_font_size > 0
                        and len(text) >= 2
                        and explicit_sizes
                        and max(explicit_sizes) < min_body_font_size
                        and not is_footer
                        and not is_header_chrome
                        and not is_decorative_index
                        and not is_compact_annotation
                    ):
                        below_requested_body_floor.append(text[:48])
                    if (
                        len(text) >= 6
                        and explicit_sizes
                        and max(explicit_sizes) <= 12.0
                        and not is_footer
                        and not is_header_chrome
                        and not is_decorative_index
                    ):
                        tiny_text_boxes.append(text[:48])
                    if explicit_sizes and likely_text_overflow(
                        getattr(shape, "text", ""), shape_width, shape_height, max(explicit_sizes)
                    ):
                        overflow_text_boxes.append(text[:48])
            if not getattr(shape, "has_text_frame", False) or not getattr(shape, "text", "").strip():
                visual_shapes += 1
            if (
                getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.AUTO_SHAPE
                and safe_auto_shape_type(shape) == MSO_SHAPE.ROUNDED_RECTANGLE
            ):
                rounded_rectangles += 1
            if bool(getattr(shape, "has_chart", False)):
                charts += 1
            if bool(getattr(shape, "has_table", False)):
                tables += 1
                if min_body_font_size > 0:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            cell_text = " ".join(cell.text.split())
                            if not cell_text:
                                continue
                            cell_sizes: list[float] = []
                            for paragraph in cell.text_frame.paragraphs:
                                if paragraph.font.size is not None:
                                    cell_sizes.append(float(paragraph.font.size.pt))
                                cell_sizes.extend(
                                    float(run.font.size.pt)
                                    for run in paragraph.runs
                                    if run.font.size is not None
                                )
                            all_explicit_sizes.extend(cell_sizes)
                            if cell_sizes and max(cell_sizes) < min_body_font_size:
                                below_requested_body_floor.append(cell_text[:48])
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1

            # A large solid shape added late in z-order can hide a complete
            # diagram/table while every underlying text box still exists. This
            # catches the common "border rectangle accidentally filled white"
            # failure that slide/shape counts cannot see.
            shape_area = max(0, shape_width) * max(0, shape_height)
            shape_fill = getattr(shape, "fill", None)
            if (
                shape_index > 0
                and shape_area >= width * height * 0.15
                and not getattr(shape, "text", "").strip()
                and getattr(shape_fill, "type", None) == MSO_FILL_TYPE.SOLID
            ):
                covered_text = 0
                right, bottom = left + shape_width, top + shape_height
                for earlier in shapes[:shape_index]:
                    if not getattr(earlier, "has_text_frame", False) or not getattr(earlier, "text", "").strip():
                        continue
                    center_x = int(earlier.left) + int(earlier.width) // 2
                    center_y = int(earlier.top) + int(earlier.height) // 2
                    if left <= center_x <= right and top <= center_y <= bottom:
                        covered_text += 1
                if covered_text >= 3:
                    errors.append(
                        f"slide {number} has a late opaque shape covering {covered_text} earlier text boxes"
                    )

        # Near-total overlap between *similarly sized* text regions is almost
        # never intentional. A large labelled panel may legitimately contain
        # smaller text boxes, so containment alone is not a collision. The
        # area-ratio guard keeps that common composition legal while still
        # catching duplicated/same-coordinate regions that file-open checks
        # miss.
        for left_index, left_box in enumerate(text_geometries):
            left_x, left_y, left_w, left_h, left_text = left_box
            left_area = left_w * left_h
            for right_box in text_geometries[left_index + 1 :]:
                right_x, right_y, right_w, right_h, right_text = right_box
                overlap_w = max(0, min(left_x + left_w, right_x + right_w) - max(left_x, right_x))
                overlap_h = max(0, min(left_y + left_h, right_y + right_h) - max(left_y, right_y))
                overlap_area = overlap_w * overlap_h
                smaller_area = min(left_area, right_w * right_h)
                larger_area = max(left_area, right_w * right_h)
                similarly_sized = smaller_area > 0 and larger_area / smaller_area <= 2.25
                if similarly_sized and overlap_area / smaller_area >= 0.80:
                    overlapping_text_pairs.append((left_text, right_text))

        if serialized_payloads:
            errors.append(
                f"slide {number} contains serialized rich-text/style payload in audience text; "
                "configuration/action dictionaries are also forbidden; use plain audience text "
                "or rich_text_box with run mappings"
            )
        if mojibake_payloads:
            errors.append(
                f"slide {number} contains likely mojibake/garbled audience text; "
                "repair the source encoding before delivery"
            )
        if below_requested_body_floor:
            errors.append(
                f"slide {number} has {len(below_requested_body_floor)} non-chrome text boxes below "
                f"the requested {min_body_font_size:g}pt body floor"
            )
        if min_title_font_size > 0:
            candidate_sizes = all_explicit_sizes if number == 1 else title_region_sizes
            authored_title_size = max(candidate_sizes, default=0.0)
            if authored_title_size < min_title_font_size:
                errors.append(
                    f"slide {number} has no title at or above the requested "
                    f"{min_title_font_size:g}pt title floor"
                )

        meaningful_chars = sum(len(text) for text in texts)
        if (
            min_body_font_size > 0
            and number > 1
            and body_text_chars < 8
            and charts + tables + pictures == 0
        ):
            errors.append(
                f"slide {number} has no substantive body content below the title region; "
                "a title and decorative chrome cannot satisfy a content slide"
            )
        if meaningful_chars < 8 and visual_shapes == 0:
            warnings.append(f"slide {number} appears blank")
        if out_of_bounds:
            errors.append(
                f"slide {number} has out-of-bounds shapes: {', '.join(out_of_bounds[:6])}"
            )
        if overlapping_text_pairs:
            examples = "; ".join(
                f"{left!r} over {right!r}"
                for left, right in overlapping_text_pairs[:3]
            )
            errors.append(
                f"slide {number} has {len(overlapping_text_pairs)} near-total overlapping text-box pairs; "
                f"assign distinct page regions before delivery: {examples}"
            )
        if meaningful_chars < 30 and visual_shapes < 2:
            warnings.append(f"slide {number} has unusually little visible content")
        if number == len(deck.slides) and len(deck.slides) > 1 and meaningful_chars < 20 and visual_shapes < 2:
            errors.append(
                f"slide {number} is an empty/generic ending; close with a decision, synthesis, or next action"
            )
        if len(tiny_text_boxes) >= 6:
            warnings.append(
                f"slide {number} has {len(tiny_text_boxes)} non-footer text boxes at 12pt or smaller; "
                "split or prioritize the page instead of shrinking the content"
            )
        if overflow_text_boxes:
            errors.append(
                f"slide {number} has {len(overflow_text_boxes)} text boxes likely to overflow at authored font size; "
                f"shorten or enlarge: {', '.join(overflow_text_boxes[:3])}"
            )
        if len(texts) >= 30:
            warnings.append(
                f"slide {number} has {len(texts)} separate text boxes; verify that the page is not a dense dashboard"
            )
        if rounded_rectangles >= 4 and rounded_rectangles * 2 >= max(1, visual_shapes):
            card_heavy_slides.append(number)
        if tables > 0:
            table_slides.append(number)
        slides.append(
            {
                "slide": number,
                "shapes": len(slide.shapes),
                "text_boxes": len(texts),
                "text_chars": meaningful_chars,
                "body_text_chars": body_text_chars,
                "visual_shapes": visual_shapes,
                "tiny_text_boxes": len(tiny_text_boxes),
                "likely_overflow_text_boxes": len(overflow_text_boxes),
                "overlapping_text_box_pairs": len(overlapping_text_pairs),
                "rounded_rectangles": rounded_rectangles,
                "charts": charts,
                "tables": tables,
                "pictures": pictures,
                "layout_shapes": sum(signature.values()),
                "layout_fingerprint": layout_fingerprint(signature),
            }
        )
    card_heavy_limit = max(3, math.ceil(len(slides) * 0.4))
    if len(card_heavy_slides) >= card_heavy_limit:
        errors.append(
            "rounded-card UI grammar dominates "
            f"{len(card_heavy_slides)} of {len(slides)} slides "
            f"({', '.join(map(str, card_heavy_slides))}); replace repeated dashboard tiles "
            "with flat editorial evidence, charts, timelines, matrices, or page-specific compositions"
        )
    table_heavy_limit = max(3, math.ceil(len(slides) * 0.3))
    if len(table_slides) >= table_heavy_limit:
        errors.append(
            "native tables dominate "
            f"{len(table_slides)} of {len(slides)} slides "
            f"({', '.join(map(str, table_slides))}); keep exact lookup in an appendix and use charts, "
            "comparisons, timelines, matrices, or page-specific explanatory compositions in the main story"
        )
    originality, originality_errors, originality_warnings = layout_originality_report(
        signatures,
        style_reference=style_reference,
        allow_repeated_layouts=allow_repeated_layouts,
    )
    errors.extend(originality_errors)
    warnings.extend(originality_warnings)
    return slides, errors, warnings, originality


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit PPTX structure, readability, and page-layout originality."
    )
    parser.add_argument("deck", type=Path, help="PPTX deck to audit")
    parser.add_argument(
        "--style-reference",
        type=Path,
        help="optional style-reference PPTX; copied page geometry becomes a hard failure",
    )
    parser.add_argument(
        "--allow-repeated-layouts",
        action="store_true",
        help="downgrade repeated body geometry to a warning for explicit recurring/fixed-schema reports",
    )
    parser.add_argument(
        "--min-body-font-size",
        type=float,
        default=0.0,
        help="hard minimum body font size in points; 0 disables the request-specific gate",
    )
    parser.add_argument(
        "--min-title-font-size",
        type=float,
        default=0.0,
        help="hard minimum slide-title font size in points; 0 disables the request-specific gate",
    )
    args = parser.parse_args()

    path = args.deck.expanduser().resolve()
    style_reference = (
        args.style_reference.expanduser().resolve() if args.style_reference is not None else None
    )
    errors: list[str] = []
    warnings: list[str] = []
    slides: list[dict[str, object]] = []
    originality: dict[str, object] = {
        "policy": "recurring" if args.allow_repeated_layouts else "audience-facing",
        "internal_repeated_pairs": [],
        "style_reference": None,
    }
    if not path.is_file():
        errors.append(f"file not found: {path}")
    else:
        try:
            errors.extend(package_errors(path))
            slide_report, presentation_errors, presentation_warnings, originality = presentation_report(
                path,
                style_reference=style_reference,
                allow_repeated_layouts=args.allow_repeated_layouts,
                min_body_font_size=max(0.0, args.min_body_font_size),
                min_title_font_size=max(0.0, args.min_title_font_size),
            )
            slides = slide_report
            errors.extend(presentation_errors)
            warnings.extend(presentation_warnings)
        except (OSError, zipfile.BadZipFile, ValueError) as exc:
            errors.append(f"cannot inspect PPTX: {exc}")

    status = "failed" if errors else ("passed_with_warnings" if warnings else "passed")
    # High-confidence composition/readability defects are promoted to errors;
    # lower-confidence density cautions remain advisory for the repair pass.
    quality_score = max(0, 100 - 25 * len(errors) - 6 * len(warnings))
    report = {
        "ok": not errors,
        "status": status,
        "path": str(path),
        "slide_count": len(slides),
        "fatal_errors": errors,
        "quality_warnings": warnings,
        "quality_score": quality_score,
        "quality_gate_passed": not errors,
        # Keep the original fields for callers that predate severity-aware output.
        "errors": errors,
        "warnings": warnings,
        "layout_originality": originality,
        "slides": slides,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors else 2


if __name__ == "__main__":
    raise SystemExit(main())
