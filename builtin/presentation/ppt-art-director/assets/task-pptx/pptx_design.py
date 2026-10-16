#!/usr/bin/env python3
"""Small, reusable 16:9 design system for task-specific ``python-pptx`` decks.

The module deliberately owns only presentation primitives.  Narrative, page
structure, facts, and brand choices belong in the task builder next to it.
Public geometry helpers accept inch scalars only and normalize them to
``float`` immediately.  ``Inches(...)``, ``Emu(...)``, and raw EMU integers
are rejected at this boundary.  Conversion to integer EMU happens exactly once,
immediately before a ``python-pptx``/OOXML API call.
"""

from __future__ import annotations

from dataclasses import dataclass
from math import isfinite
from numbers import Integral
import platform
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Length, Pt


EMU_PER_INCH = 914_400
SLIDE_W = 12_192_000
SLIDE_H = 6_858_000
SLIDE_W_IN = SLIDE_W / EMU_PER_INCH
SLIDE_H_IN = SLIDE_H / EMU_PER_INCH
MODEL_GEOMETRY_TOLERANCE_IN = 0.1
# Normal slide geometry expressed in inches is far below this value. Large
# integers at a public helper boundary are almost certainly raw EMU values.
EMU_LIKE_INTEGER_THRESHOLD = 1_000


def platform_font_pair() -> tuple[str, str]:
    """Return a native CJK/Latin pair for the current generation host.

    The deck renderer is intentionally pure ``python-pptx`` on every platform,
    but font metrics still affect wrapping during native/LibreOffice preview.
    Use fonts shipped with each desktop family instead of writing a Windows-only
    face into decks generated on macOS or Linux.  PowerPoint keeps the text
    editable and performs its normal fallback when the file is opened elsewhere.
    """

    system = platform.system().lower()
    if system == "darwin":
        return "PingFang SC", "Helvetica Neue"
    if system == "windows":
        return "Microsoft YaHei", "Arial"
    return "Noto Sans CJK SC", "Liberation Sans"


PLATFORM_FONT_CN, PLATFORM_FONT_EN = platform_font_pair()


@dataclass(frozen=True)
class Theme:
    """Regular custom baseline: restrained consulting editorial system."""

    navy: RGBColor = RGBColor(0x05, 0x1C, 0x2C)
    ink: RGBColor = RGBColor(0x23, 0x2B, 0x31)
    muted: RGBColor = RGBColor(0x58, 0x5E, 0x63)
    line: RGBColor = RGBColor(0xD8, 0xD9, 0xD6)
    paper: RGBColor = RGBColor(0xF5, 0xF3, 0xEE)
    white: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)
    accent: RGBColor = RGBColor(0xBF, 0xA9, 0x60)
    accent_soft: RGBColor = RGBColor(0xEE, 0xE8, 0xD3)
    positive: RGBColor = RGBColor(0x4E, 0x8C, 0x4D)
    positive_soft: RGBColor = RGBColor(0xE2, 0xEE, 0xDF)
    warning: RGBColor = RGBColor(0xD8, 0x97, 0x20)
    warning_soft: RGBColor = RGBColor(0xFB, 0xF0, 0xD6)
    negative: RGBColor = RGBColor(0xB3, 0x4A, 0x42)
    negative_soft: RGBColor = RGBColor(0xF5, 0xE3, 0xE0)
    font_cn: str = PLATFORM_FONT_CN
    font_en: str = PLATFORM_FONT_EN


DEFAULT_THEME = Theme()


@dataclass(frozen=True)
class LayoutRegion:
    """A content region in inches, with no decorative panel or implied styling."""

    x: float
    y: float
    w: float
    h: float

    @property
    def box(self) -> tuple[float, float, float, float]:
        return self.x, self.y, self.w, self.h


def layout_regions(
    family: str,
    *,
    x: float = 0.62,
    y: float = 1.72,
    w: float = 12.09,
    h: float = 5.06,
    gap: float = 0.48,
    primary_share: float = 0.68,
    mirror: bool = False,
) -> dict[str, LayoutRegion]:
    """Compose evidence and interpretation without another grid of cards.

    editorial: a dominant exhibit with a narrower explanation column.
    spotlight: a wide exhibit above a compact interpretation strip.
    comparison: two equal exhibits, with one shared conclusion underneath.
    metric: a narrow typographic anchor beside a larger explanation/diagram.

    Regions contain no facts and draw no shapes. The task renderer owns the
    chart, photo, diagram or copy placed in each region. Title/footer stay outside
    the default body. Mirror changes reading direction without changing scale.
    """
    import math

    values = (x, y, w, h, gap, primary_share)
    if type(mirror) is not bool:
        raise ValueError("mirror must be boolean")
    if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
        raise ValueError("layout geometry must contain finite numbers")
    if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > SLIDE_W_IN + 0.001 or y + h > SLIDE_H_IN + 0.001:
        raise ValueError("layout region must fit inside the slide")
    if gap < 0.15 or gap >= min(w, h) / 3 or not 0.55 <= primary_share <= 0.78:
        raise ValueError("use a readable gap and a primary_share within 0.55..0.78")
    if family == "editorial":
        pw = (w - gap) * primary_share
        regions = {"primary": LayoutRegion(x, y, pw, h), "secondary": LayoutRegion(x + pw + gap, y, w - gap - pw, h)}
    elif family == "spotlight":
        ph = (h - gap) * primary_share
        regions = {"primary": LayoutRegion(x, y, w, ph), "secondary": LayoutRegion(x, y + ph + gap, w, h - gap - ph)}
    elif family == "comparison":
        pw = (w - gap) / 2
        ph = (h - gap) * 0.76
        regions = {"primary": LayoutRegion(x, y, pw, ph), "secondary": LayoutRegion(x + pw + gap, y, pw, ph), "conclusion": LayoutRegion(x, y + ph + gap, w, h - gap - ph)}
    elif family == "metric":
        pw = (w - gap) * (1 - primary_share)
        regions = {"primary": LayoutRegion(x, y, pw, h), "secondary": LayoutRegion(x + pw + gap, y, w - gap - pw, h)}
    else:
        raise ValueError("layout family must be editorial, spotlight, comparison, or metric")
    if mirror:
        regions = {name: LayoutRegion(x + w - (r.x - x) - r.w, r.y, r.w, r.h) for name, r in regions.items()}
    return regions

# Airy, light product-story baseline inspired by Google's editorial and
# Material visual language. It uses blue for hierarchy and reserves the other
# brand-family colors for semantic meaning instead of rainbow decoration.
GOOGLE_CLEAN_THEME = Theme(
    navy=RGBColor(0x42, 0x85, 0xF4),
    ink=RGBColor(0x20, 0x21, 0x24),
    muted=RGBColor(0x5F, 0x63, 0x68),
    line=RGBColor(0xDA, 0xDC, 0xE0),
    paper=RGBColor(0xF8, 0xF9, 0xFA),
    white=RGBColor(0xFF, 0xFF, 0xFF),
    accent=RGBColor(0x42, 0x85, 0xF4),
    accent_soft=RGBColor(0xE8, 0xF0, 0xFE),
    positive=RGBColor(0x34, 0xA8, 0x53),
    positive_soft=RGBColor(0xE6, 0xF4, 0xEA),
    warning=RGBColor(0xFB, 0xBC, 0x04),
    warning_soft=RGBColor(0xFE, 0xF7, 0xE0),
    negative=RGBColor(0xEA, 0x43, 0x35),
    negative_soft=RGBColor(0xFC, 0xE8, 0xE6),
    font_cn=PLATFORM_FONT_CN,
    font_en=PLATFORM_FONT_EN,
)


DEFAULT_SEMANTIC_LEGEND = (
    ("事实", RGBColor(0x62, 0x63, 0x5E)),
    ("计算", RGBColor(0x78, 0x5B, 0x72)),
    ("建议", RGBColor(0x3C, 0x75, 0x65)),
)


@dataclass(frozen=True)
class Edge:
    """Connector geometry used by :func:`edges_first`."""

    start: tuple[float, float]
    end: tuple[float, float]
    color: RGBColor | str | tuple[int, int, int] = DEFAULT_THEME.line
    weight: float = 1.5
    dash: bool = False
    arrow: bool = True


def _as_inches(value: float | int) -> float:
    """Normalize one public geometry value to an inch ``float``."""

    if isinstance(value, bool):
        raise TypeError("boolean is not a valid PowerPoint length")
    if isinstance(value, Length):
        raise TypeError(
            "geometry expects an inch float, not Inches()/Emu(); "
            "pass the original inch value"
        )
    if isinstance(value, Integral) and abs(int(value)) >= EMU_LIKE_INTEGER_THRESHOLD:
        raise TypeError(
            "geometry expects an inch float, not a raw EMU integer; "
            "divide by EMU_PER_INCH before calling the helper"
        )
    try:
        inches = float(value)
    except (TypeError, ValueError) as exc:
        raise TypeError(f"unsupported PowerPoint length: {value!r}") from exc
    if not isfinite(inches):
        raise ValueError(f"PowerPoint length must be finite: {value!r}")
    return inches


def i_emu(value: float | int) -> int:
    """Convert an inch scalar to an integer EMU value at the API boundary."""

    return int(round(_as_inches(value) * EMU_PER_INCH))


def _rgb(value: RGBColor | str | tuple[int, int, int]) -> RGBColor:
    if isinstance(value, RGBColor):
        return value
    if isinstance(value, str):
        compact = value.strip().lstrip("#")
        if len(compact) != 6:
            raise ValueError(f"expected a six-digit RGB color, got {value!r}")
        return RGBColor.from_string(compact.upper())
    if isinstance(value, tuple) and len(value) == 3:
        return RGBColor(*(int(channel) for channel in value))
    raise TypeError(f"unsupported color value: {value!r}")


def _vertical_anchor(value):
    """Return a vertical-anchor member owned by this python-pptx runtime.

    python-pptx 0.6.x used integer-like ``EnumValue`` objects while 1.x uses a
    stdlib-style ``MSO_VERTICAL_ANCHOR`` enum.  Checking ``isinstance`` against
    the enum therefore breaks across versions and can also reject an equivalent
    member imported through the older enum name.  Normalize by semantic member
    name/value and always return the member from the currently loaded runtime.
    """

    if value is None:
        return None

    aliases = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "center": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }

    # The common path also supports legacy python-pptx EnumValue instances,
    # which compare equal to the local integer-backed enum members.
    for member in (MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM):
        try:
            if value == member:
                return member
        except (TypeError, ValueError):
            pass

    member_name = getattr(value, "name", None)
    if isinstance(member_name, str):
        resolved = aliases.get(member_name.strip().lower())
        if resolved is not None:
            return resolved

    if isinstance(value, str):
        # Accept model-friendly aliases plus qualified names sometimes emitted
        # by generated task code, e.g. "MSO_VERTICAL_ANCHOR.MIDDLE".
        key = value.strip().rsplit(".", 1)[-1].split(" ", 1)[0].lower()
        resolved = aliases.get(key)
        if resolved is not None:
            return resolved

    numeric_value = getattr(value, "value", value)
    if isinstance(numeric_value, Integral) and not isinstance(numeric_value, bool):
        numeric_aliases = {
            int(MSO_ANCHOR.TOP): MSO_ANCHOR.TOP,
            int(MSO_ANCHOR.MIDDLE): MSO_ANCHOR.MIDDLE,
            int(MSO_ANCHOR.BOTTOM): MSO_ANCHOR.BOTTOM,
        }
        resolved = numeric_aliases.get(int(numeric_value))
        if resolved is not None:
            return resolved

    if isinstance(value, type):
        raise ValueError(
            f"vertical anchor received enum class {value.__name__}; "
            "pass a member such as MSO_ANCHOR.MIDDLE"
        )
    raise ValueError(
        f"unsupported vertical anchor {value!r}; use MSO_ANCHOR.TOP, "
        "MSO_ANCHOR.MIDDLE, or MSO_ANCHOR.BOTTOM"
    )


def _box_emu(x: float, y: float, w: float, h: float) -> tuple[int, int, int, int]:
    x, y, w, h = (_as_inches(value) for value in (x, y, w, h))
    right_overshoot = x + w - SLIDE_W_IN
    bottom_overshoot = y + h - SLIDE_H_IN
    if 0 < right_overshoot <= MODEL_GEOMETRY_TOLERANCE_IN:
        w = SLIDE_W_IN - x
    if 0 < bottom_overshoot <= MODEL_GEOMETRY_TOLERANCE_IN:
        h = SLIDE_H_IN - y
    left, top, width, height = i_emu(x), i_emu(y), i_emu(w), i_emu(h)
    if width < 0 or height < 0:
        raise ValueError(f"negative box size: {(x, y, w, h)!r}")
    if left < 0 or top < 0 or left + width > SLIDE_W or top + height > SLIDE_H:
        raise ValueError(f"box is outside the 16:9 canvas: {(x, y, w, h)!r}")
    return left, top, width, height


def _point_emu(x: float, y: float) -> tuple[int, int]:
    x, y = _as_inches(x), _as_inches(y)
    if 0 < x - SLIDE_W_IN <= MODEL_GEOMETRY_TOLERANCE_IN:
        x = SLIDE_W_IN
    if 0 < y - SLIDE_H_IN <= MODEL_GEOMETRY_TOLERANCE_IN:
        y = SLIDE_H_IN
    point = i_emu(x), i_emu(y)
    if point[0] < 0 or point[1] < 0 or point[0] > SLIDE_W or point[1] > SLIDE_H:
        raise ValueError(f"point is outside the 16:9 canvas: {(x, y)!r}")
    return point


def _set_fill(shape, fill: RGBColor | str | tuple[int, int, int] | None) -> None:
    """Apply a solid fill, or a true no-fill when ``fill`` is ``None``."""

    if fill is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)


def _set_line(
    shape,
    color: RGBColor | str | tuple[int, int, int] | None,
    width: float | None = None,
) -> None:
    if color is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = _rgb(color)
    if width is not None:
        shape.line.width = Pt(float(width))


def _disable_shadow(shape) -> None:
    try:
        shape.shadow.inherit = False
    except (AttributeError, ValueError):
        pass


def _send_to_back(shape) -> None:
    tree = shape._element.getparent()
    tree.remove(shape._element)
    tree.insert(2, shape._element)


def new_presentation() -> Presentation:
    deck = Presentation()
    deck.slide_width = Emu(SLIDE_W)
    deck.slide_height = Emu(SLIDE_H)
    return deck


def new_slide(deck: Presentation, background=DEFAULT_THEME.white):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_background(slide, background)
    return slide


def add_background(slide, fill=DEFAULT_THEME.white):
    shape = rect(slide, 0, 0, SLIDE_W_IN, SLIDE_H_IN, fill=fill)
    _send_to_back(shape)
    return shape


def image_cover(
    slide,
    path: str | Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    focal_x: float = 0.5,
    focal_y: float = 0.5,
):
    """Place a real raster image with a focal-point-aware cover crop.

    The picture remains a native PowerPoint picture object. ``focal_x`` and
    ``focal_y`` are normalized 0..1 coordinates in the source image and shift
    the crop toward the subject without stretching the pixels.
    """

    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"image asset does not exist: {source}")
    focal_x = max(0.0, min(1.0, float(focal_x)))
    focal_y = max(0.0, min(1.0, float(focal_y)))
    with Image.open(source) as image:
        image_w, image_h = image.size
    if image_w <= 0 or image_h <= 0:
        raise ValueError(f"image asset has invalid dimensions: {source}")

    left, top, width, height = _box_emu(x, y, w, h)
    picture = slide.shapes.add_picture(str(source), left, top, width, height)
    source_ratio = image_w / image_h
    target_ratio = _as_inches(w) / _as_inches(h)
    if source_ratio > target_ratio:
        visible = target_ratio / source_ratio
        room = 1.0 - visible
        crop_left = max(0.0, min(room, focal_x - visible / 2.0))
        picture.crop_left = crop_left
        picture.crop_right = room - crop_left
    elif source_ratio < target_ratio:
        visible = source_ratio / target_ratio
        room = 1.0 - visible
        crop_top = max(0.0, min(room, focal_y - visible / 2.0))
        picture.crop_top = crop_top
        picture.crop_bottom = room - crop_top
    return picture


def image_contain(
    slide,
    path: str | Path,
    x: float,
    y: float,
    w: float,
    h: float,
):
    """Place a real raster image inside a box without cropping or stretching."""

    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"image asset does not exist: {source}")
    with Image.open(source) as image:
        image_w, image_h = image.size
    if image_w <= 0 or image_h <= 0:
        raise ValueError(f"image asset has invalid dimensions: {source}")
    target_w, target_h = _as_inches(w), _as_inches(h)
    scale = min(target_w / image_w, target_h / image_h)
    placed_w, placed_h = image_w * scale, image_h * scale
    placed_x = _as_inches(x) + (target_w - placed_w) / 2.0
    placed_y = _as_inches(y) + (target_h - placed_h) / 2.0
    return slide.shapes.add_picture(
        str(source),
        i_emu(placed_x),
        i_emu(placed_y),
        i_emu(placed_w),
        i_emu(placed_h),
    )


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill=DEFAULT_THEME.white,
    line=None,
    line_width: float | None = None,
    theme: Theme | None = None,
):
    left, top, width, height = _box_emu(x, y, w, h)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _set_fill(shape, fill)
    _set_line(shape, line, line_width)
    _disable_shadow(shape)
    return shape


def round_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *legacy_corner,
    fill=DEFAULT_THEME.white,
    line=None,
    line_width: float | None = None,
    corner: float = 0.12,
    color=None,
    radius: float | None = None,
    theme: Theme | None = None,
):
    if legacy_corner:
        if len(legacy_corner) != 1:
            raise TypeError("round_rect accepts at most one legacy positional corner value")
        corner = float(legacy_corner[0])
    if radius is not None:
        corner = float(radius)
    if line is None and color is not None:
        line = color
    left, top, width, height = _box_emu(x, y, w, h)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.adjustments[0] = max(0.0, min(float(corner), 0.5))
    _set_fill(shape, fill)
    _set_line(shape, line, line_width)
    _disable_shadow(shape)
    return shape


def circle(
    slide,
    x: float,
    y: float,
    diameter: float,
    *legacy_height,
    fill=DEFAULT_THEME.accent,
    line=None,
    line_width: float | None = None,
    theme: Theme | None = None,
):
    height = diameter
    if legacy_height:
        if len(legacy_height) != 1:
            raise TypeError("circle accepts at most one legacy height value")
        height = legacy_height[0]
    left, top, width, height = _box_emu(x, y, diameter, height)
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    _set_fill(shape, fill)
    _set_line(shape, line, line_width)
    _disable_shadow(shape)
    return shape


def line(
    slide,
    x1,
    y1=None,
    x2=None,
    y2=None,
    *,
    color=DEFAULT_THEME.line,
    weight: float = 1.0,
    dash: bool = False,
    arrow: bool = False,
    theme: Theme | None = None,
):
    if isinstance(x1, (tuple, list)) and isinstance(y1, (tuple, list)) and x2 is None and y2 is None:
        (x1, y1), (x2, y2) = x1, y1
    if y1 is None or x2 is None or y2 is None:
        raise TypeError("line requires x1, y1, x2, y2 or two point tuples")
    start_x, start_y = _point_emu(x1, y1)
    end_x, end_y = _point_emu(x2, y2)
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y
    )
    connector.line.color.rgb = _rgb(color)
    connector.line.width = Pt(float(weight))
    _disable_shadow(connector)
    line_xml = connector.line._get_or_add_ln()
    if dash:
        dash_xml = etree.SubElement(line_xml, qn("a:prstDash"))
        dash_xml.set("val", "dash")
    if arrow:
        # DrawingML names the destination end of a line ``tailEnd``.
        arrow_xml = etree.SubElement(line_xml, qn("a:tailEnd"))
        arrow_xml.set("type", "triangle")
        arrow_xml.set("w", "sm")
        arrow_xml.set("len", "sm")
    return connector


def edges_first(slide, edges: Iterable[Edge | Mapping[str, object]]):
    """Create all connectors before the caller creates any node/card shapes.

    This explicit two-phase API keeps connectors behind nodes and labels.  Pass
    :class:`Edge` objects or dictionaries accepted by the ``Edge`` constructor.
    """

    connectors = []
    for raw in edges:
        edge = raw if isinstance(raw, Edge) else Edge(**raw)
        connectors.append(
            line(
                slide,
                edge.start[0],
                edge.start[1],
                edge.end[0],
                edge.end[1],
                color=edge.color,
                weight=edge.weight,
                dash=edge.dash,
                arrow=edge.arrow,
            )
        )
    return connectors


def _set_font_family(font, font_cn: str, font_en: str) -> None:
    """Set both Latin and East Asian typefaces, including chart labels."""
    font.name = font_en
    properties = font._element
    east_asia = properties.find(qn("a:ea"))
    if east_asia is None:
        # Use DrawingML's font-child order (latin, ea, cs, sym).
        east_asia = etree.Element(qn("a:ea"))
        latin = properties.find(qn("a:latin"))
        properties.insert(properties.index(latin) + 1, east_asia)
    east_asia.set("typeface", font_cn)


def _apply_run_font(
    run,
    *,
    size: float,
    bold: bool,
    italic: bool = False,
    color,
    font_cn: str,
    font_en: str,
) -> None:
    _set_font_family(run.font, font_cn, font_en)
    run.font.size = Pt(float(size))
    run.font.bold = bool(bold)
    run.font.italic = bool(italic)
    run.font.color.rgb = _rgb(color)


def text_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str | Sequence[Mapping[str, object]],
    *,
    size: float = 16,
    font_size: float | None = None,
    bold: bool = False,
    italic: bool = False,
    color=DEFAULT_THEME.ink,
    font_cn: str = DEFAULT_THEME.font_cn,
    font_en: str = DEFAULT_THEME.font_en,
    font: str | None = None,
    theme: Theme | None = None,
    fill=None,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    margins: float = 0,
    wrap: bool = True,
):
    if font_size is not None:
        size = float(font_size)
    if font is not None:
        font_cn = str(font)
        font_en = str(font)
    if theme is not None and color == DEFAULT_THEME.ink:
        color = theme.ink
    rich_runs: list[Mapping[str, object]] | None = None
    if isinstance(text, Sequence) and not isinstance(text, (str, bytes, bytearray)):
        rich_runs = list(text)
        if not all(isinstance(item, Mapping) for item in rich_runs):
            raise TypeError("text_box rich text must be a sequence of mappings")
        if anchor == MSO_ANCHOR.TOP and rich_runs and "anchor" in rich_runs[0]:
            anchor = rich_runs[0]["anchor"]

    left, top, width, height = _box_emu(x, y, w, h)
    box = slide.shapes.add_textbox(left, top, width, height)
    _set_fill(box, fill)
    frame = box.text_frame
    margin = Emu(i_emu(margins))
    frame.margin_left = margin
    frame.margin_right = margin
    frame.margin_top = margin
    frame.margin_bottom = margin
    frame.word_wrap = bool(wrap)
    frame.vertical_anchor = _vertical_anchor(anchor)
    if rich_runs is None:
        for index, value in enumerate(str(text).split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            run = paragraph.add_run()
            run.text = value
            _apply_run_font(
                run,
                size=size,
                bold=bold,
                italic=italic,
                color=color,
                font_cn=font_cn,
                font_en=font_en,
            )
    else:
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        for item_index, item in enumerate(rich_runs):
            values = str(item.get("text", "")).split("\n")
            for value_index, value in enumerate(values):
                if value_index > 0:
                    paragraph = frame.add_paragraph()
                if item_index == 0 or value_index > 0:
                    paragraph.alignment = item.get("align", align)
                run = paragraph.add_run()
                run.text = value
                _apply_run_font(
                    run,
                    size=float(item.get("size", size)),
                    bold=bool(item.get("bold", bold)),
                    italic=bool(item.get("italic", italic)),
                    color=item.get("color", color),
                    font_cn=str(item.get("font_cn", font_cn)),
                    font_en=str(item.get("font_en", font_en)),
                )
    return box


def rich_text_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    runs: Sequence[Mapping[str, object]],
    *,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    theme: Theme = DEFAULT_THEME,
):
    left, top, width, height = _box_emu(x, y, w, h)
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.margin_left = frame.margin_right = Emu(0)
    frame.margin_top = frame.margin_bottom = Emu(0)
    frame.word_wrap = True
    frame.vertical_anchor = _vertical_anchor(anchor)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for item in runs:
        run = paragraph.add_run()
        run.text = str(item.get("text", ""))
        _apply_run_font(
            run,
            size=float(item.get("size", 16)),
            bold=bool(item.get("bold", False)),
            italic=bool(item.get("italic", False)),
            color=item.get("color", theme.ink),
            font_cn=str(item.get("font_cn", theme.font_cn)),
            font_en=str(item.get("font_en", theme.font_en)),
        )
    return box


def header_footer(
    slide,
    *,
    brand: str,
    title: str = "",
    section: str,
    page: int,
    total: int,
    theme: Theme = DEFAULT_THEME,
    dark_mode: bool = False,
) -> None:
    """Draw the stable editorial frame used by every content slide.

    The default is a quiet light evidence page: eyebrow, conclusion headline,
    hairline, and compact footer. ``dark_mode=True`` is reserved for deliberate
    contrast anchors; it is not the whole-deck default.
    """

    title = str(title or "").strip()
    title_size = 40 if len(title) <= 14 else 37 if len(title) <= 24 else 35

    if dark_mode:
        rect(slide, 0, 0, SLIDE_W_IN, 1.24, fill=theme.navy)
        rect(slide, 0, 1.24, SLIDE_W_IN, 0.055, fill=theme.accent)
        text_box(
            slide, 0.58, 0.35, 10.15, 0.72, title,
            size=title_size, bold=True, color=theme.white,
            anchor=MSO_ANCHOR.MIDDLE, wrap=True,
        )
        text_box(
            slide, 9.82, 0.23, 1.78, 0.32, section,
            size=10, bold=True, color=theme.accent,
            align=PP_ALIGN.RIGHT, wrap=False,
        )
        text_box(
            slide, 11.93, 0.23, 0.82, 0.32, f"{page:02d} / {total:02d}",
            size=9, color=theme.white, align=PP_ALIGN.RIGHT, wrap=False,
        )
        rect(slide, 0, 7.18, SLIDE_W_IN, 0.32, fill=theme.navy)
        text_box(slide, 0.58, 7.245, 8.0, 0.15, brand, size=8, bold=True, color=theme.white, wrap=False)
        text_box(
            slide, 11.8, 7.245, 0.95, 0.15, f"P.{page:02d}",
            size=8, color=theme.accent, align=PP_ALIGN.RIGHT, wrap=False,
        )
        return

    text_box(
        slide, 0.58, 0.12, 7.8, 0.24, section.upper(),
        size=10, bold=True, color=theme.accent, wrap=False,
    )
    text_box(
        slide, 11.25, 0.12, 1.5, 0.24, f"{page:02d} / {total:02d}",
        size=9, color=theme.muted, align=PP_ALIGN.RIGHT, wrap=False,
    )
    text_box(
        slide, 0.58, 0.40, 12.17, 0.78, title,
        size=title_size, bold=True, color=theme.ink,
        anchor=MSO_ANCHOR.MIDDLE, wrap=True,
    )
    line(slide, 0.58, 1.30, 12.75, 1.30, color=theme.line, weight=0.8)
    rect(slide, 0.58, 1.285, 0.7, 0.035, fill=theme.accent)
    rect(slide, 0, 7.18, SLIDE_W_IN, 0.32, fill=theme.paper)
    line(slide, 0, 7.18, SLIDE_W_IN, 7.18, color=theme.line, weight=0.6)
    text_box(slide, 0.58, 7.245, 8.0, 0.15, brand, size=8, color=theme.muted, wrap=False)
    text_box(
        slide, 11.8, 7.245, 0.95, 0.15, f"P.{page:02d}",
        size=8, bold=True, color=theme.accent,
        align=PP_ALIGN.RIGHT, wrap=False,
    )


def semantic_legend(
    slide,
    x: float,
    y: float,
    *,
    items: Sequence[
        tuple[str, RGBColor | str | tuple[int, int, int]]
    ] = DEFAULT_SEMANTIC_LEGEND,
    item_width: float = 0.9,
    size: float = 9,
    theme: Theme = DEFAULT_THEME,
) -> float:
    """Draw a compact evidence-role legend and return the next free x position."""

    cursor = _as_inches(x)
    y = _as_inches(y)
    item_width = _as_inches(item_width)
    for label, color in items:
        rect(slide, cursor, y + 0.05, 0.12, 0.12, fill=color)
        text_box(
            slide,
            cursor + 0.17,
            y,
            max(0.2, item_width - 0.17),
            0.22,
            label,
            size=size,
            color=theme.muted,
            font_cn=theme.font_cn,
            font_en=theme.font_en,
            anchor=MSO_ANCHOR.MIDDLE,
            wrap=False,
        )
        cursor += item_width
    return cursor


def decision_frame(
    slide,
    *,
    deck_title: str,
    takeaway: str,
    brand: str,
    page: int,
    total: int,
    source_note: str = "",
    legend: Sequence[
        tuple[str, RGBColor | str | tuple[int, int, int]]
    ] = DEFAULT_SEMANTIC_LEGEND,
    theme: Theme = DEFAULT_THEME,
) -> None:
    """Create the claim-led frame used by dense executive decision pages.

    The frame owns hierarchy and evidence semantics, not page composition.  A
    task builder remains responsible for the unique visual body between the
    double rule at 1.53 inches and the footer rail at 7.18 inches.
    """

    deck_title = str(deck_title or "").strip()
    deck_title_size = 40 if len(deck_title) <= 14 else 37 if len(deck_title) <= 24 else 35

    rect(slide, 0, 0, SLIDE_W_IN, 0.32, fill=theme.navy)
    rect(slide, 12.4, 0, SLIDE_W_IN - 12.4, 0.32, fill=theme.accent)
    text_box(
        slide,
        0.5,
        0.04,
        7.4,
        0.22,
        brand,
        size=9,
        bold=True,
        color=theme.white,
        font_cn=theme.font_cn,
        font_en=theme.font_en,
        anchor=MSO_ANCHOR.MIDDLE,
        wrap=False,
    )
    text_box(
        slide,
        8.1,
        0.04,
        4.35,
        0.22,
        f"P.{page:02d} / {total:02d}",
        size=9,
        color=theme.white,
        font_cn=theme.font_cn,
        font_en=theme.font_en,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
        wrap=False,
    )
    text_box(
        slide,
        0.5,
        0.5,
        12.3,
        0.58,
        deck_title,
        size=deck_title_size,
        bold=True,
        color=theme.navy,
        font_cn=theme.font_cn,
        font_en=theme.font_en,
        wrap=True,
    )
    text_box(
        slide,
        0.5,
        1.03,
        12.3,
        0.34,
        takeaway,
        size=16,
        color=theme.ink,
        font_cn=theme.font_cn,
        font_en=theme.font_en,
        wrap=False,
    )
    line(slide, 0.5, 1.50, 12.83, 1.50, color=theme.navy, weight=1.25)
    line(slide, 0.5, 1.53, 12.83, 1.53, color=theme.accent, weight=0.5)
    rect(slide, 0, 7.18, SLIDE_W_IN, SLIDE_H_IN - 7.18, fill=theme.paper)
    if legend:
        semantic_legend(slide, 0.5, 7.23, items=legend, theme=theme)
    if source_note:
        text_box(
            slide,
            7.0,
            7.23,
            5.8,
            0.22,
            source_note,
            size=9,
            color=theme.muted,
            font_cn=theme.font_cn,
            font_en=theme.font_en,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
            wrap=False,
        )


def proportional_bar(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    value: float,
    max_value: float,
    *,
    color=DEFAULT_THEME.navy,
    track=DEFAULT_THEME.paper,
):
    """Draw a bounded comparison bar and return ``(track, fill, ratio)``."""

    maximum = float(max_value)
    if maximum <= 0:
        raise ValueError("max_value must be greater than zero")
    x, y, w, h = (_as_inches(item) for item in (x, y, w, h))
    ratio = max(0.0, min(float(value) / maximum, 1.0))
    track_shape = rect(slide, x, y, w, h, fill=track)
    fill_shape = None
    if ratio > 0:
        fill_shape = rect(slide, x, y, w * ratio, h, fill=color)
    return track_shape, fill_shape, ratio


def chip(
    slide,
    x: float,
    y: float,
    *args,
    text: str | None = None,
    width: float | None = None,
    height: float = 0.36,
    fill=DEFAULT_THEME.accent_soft,
    color=DEFAULT_THEME.accent,
    theme: Theme = DEFAULT_THEME,
    size: float = 10,
    bold: bool = True,
):
    if text is None:
        if not args:
            text = ""
        elif len(args) == 2 and all(isinstance(value, (int, float)) for value in args):
            width = _as_inches(args[0])
            height = _as_inches(args[1])
            text = ""
        else:
            text = str(args[0])
            if len(args) == 2:
                height = _as_inches(args[1])
            elif len(args) > 2:
                raise TypeError("chip canonical form accepts text and at most one positional height")
    else:
        if len(args) >= 1:
            width = _as_inches(args[0])
        if len(args) >= 2:
            height = _as_inches(args[1])
        if len(args) > 2:
            raise TypeError("chip legacy form accepts positional width and height only")
    if width is None:
        visual_units = sum(2 if ord(char) > 255 else 1 for char in str(text))
        width = max(0.78, min(3.5, 0.28 + visual_units * 0.085))
    x, y, width, height = (
        _as_inches(item) for item in (x, y, width, height)
    )
    shape = round_rect(slide, x, y, width, height, fill=fill, corner=0.25)
    text_box(
        slide,
        x + 0.12,
        y + 0.06,
        width - 0.24,
        height - 0.1,
        text,
        size=size,
        bold=bold,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        font_cn=theme.font_cn,
        font_en=theme.font_en,
        wrap=False,
    )
    return shape


def card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    title: str = "",
    body: str = "",
    accent=DEFAULT_THEME.accent,
    fill=DEFAULT_THEME.paper,
    line_color=None,
    line=None,
    line_width: float | None = None,
    radius: float | None = None,
    title_size: float = 24,
    body_size: float = 16,
    theme: Theme = DEFAULT_THEME,
):
    """Draw a restrained editorial card and return its background shape."""

    x, y, w, h = (_as_inches(item) for item in (x, y, w, h))
    if line_color is None and line is not None:
        line_color = line
    corner = 0.0 if radius is None else float(radius)
    if corner > 0:
        background = round_rect(
            slide,
            x,
            y,
            w,
            h,
            fill=fill,
            line=line_color,
            line_width=0.75 if line_width is None else line_width,
            corner=corner,
        )
    else:
        background = rect(
            slide,
            x,
            y,
            w,
            h,
            fill=fill,
            line=line_color,
            line_width=0.75 if line_width is None else line_width,
        )
    rect(slide, x, y, 0.08, h, fill=accent)
    if title:
        text_box(
            slide,
            x + 0.3,
            y + 0.24,
            w - 0.55,
            0.42,
            title,
            size=title_size,
            bold=True,
            color=theme.ink,
            font_cn=theme.font_cn,
            font_en=theme.font_en,
        )
    if body:
        text_box(
            slide,
            x + 0.3,
            y + 0.82,
            w - 0.55,
            max(0.24, h - 1.05),
            body,
            size=body_size,
            color=theme.muted,
            font_cn=theme.font_cn,
            font_en=theme.font_en,
        )
    return background


def _theme_palette(
    theme: Theme,
    colors: Sequence[RGBColor | str | tuple[int, int, int]] | None,
) -> tuple[RGBColor | str | tuple[int, int, int], ...]:
    palette = tuple(colors) if colors is not None else (
        theme.accent,
        theme.positive,
        theme.navy,
        theme.muted,
    )
    if not palette:
        raise ValueError("colors must contain at least one color")
    return palette


def native_chart(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    categories: Sequence[object],
    series: Sequence[Mapping[str, object]],
    *,
    chart_type: str = "column",
    title: str = "",
    theme: Theme = DEFAULT_THEME,
    series_colors: Sequence[RGBColor | str | tuple[int, int, int]] | None = None,
    title_size: float = 18,
    label_size: float = 12,
    title_color: RGBColor | str | tuple[int, int, int] | None = None,
    label_color: RGBColor | str | tuple[int, int, int] | None = None,
    grid_color: RGBColor | str | tuple[int, int, int] | None = None,
    show_legend: bool = False,
    show_data_labels: bool = False,
):
    """Add a native, editable chart from explicit categories and series."""
    normalized_chart_type = str(chart_type or "column").strip().lower().replace("_", "-").replace(" ", "-")
    chart_aliases = {
        "clustered-column": "column",
        "clustered-bar": "bar",
        "line-markers": "line",
        "stacked-bar": "bar-stacked",
        "bar-stack": "bar-stacked",
        "stacked-column": "column-stacked",
        "column-stack": "column-stacked",
    }
    normalized_chart_type = chart_aliases.get(normalized_chart_type, normalized_chart_type)
    chart_types = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column-stacked": XL_CHART_TYPE.COLUMN_STACKED,
        "bar-stacked": XL_CHART_TYPE.BAR_STACKED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }
    if normalized_chart_type not in chart_types:
        raise ValueError(f"unsupported chart_type: {chart_type!r}")
    categories = list(categories)
    series = list(series)
    if not categories or not series:
        raise ValueError("native_chart requires categories and series")

    data = CategoryChartData()
    data.categories = [str(value) for value in categories]
    for index, item in enumerate(series):
        if "name" not in item or "values" not in item:
            raise ValueError(f"series {index} must contain explicit name and values")
        values = list(item["values"])
        if len(values) != len(categories):
            raise ValueError(f"series {index} length must match categories")
        data.add_series(str(item["name"]), values)

    left, top, width, height = _box_emu(x, y, w, h)
    chart = slide.shapes.add_chart(chart_types[normalized_chart_type], left, top, width, height, data).chart
    palette = _theme_palette(theme, series_colors)
    resolved_title_color = theme.ink if title_color is None else title_color
    resolved_label_color = theme.muted if label_color is None else label_color
    resolved_grid_color = theme.line if grid_color is None else grid_color
    _set_font_family(chart.font, theme.font_cn, theme.font_en)
    chart.font.size = Pt(float(label_size))
    chart.font.color.rgb = _rgb(resolved_label_color)
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = str(title)
        for paragraph in chart.chart_title.text_frame.paragraphs:
            for run in paragraph.runs:
                _apply_run_font(
                    run,
                    size=title_size,
                    bold=True,
                    color=resolved_title_color,
                    font_cn=theme.font_cn,
                    font_en=theme.font_en,
                )
    chart.has_legend = bool(show_legend)
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        _set_font_family(chart.legend.font, theme.font_cn, theme.font_en)
        chart.legend.font.size = Pt(float(label_size))
        chart.legend.font.color.rgb = _rgb(resolved_label_color)
    for index, chart_series in enumerate(chart.series):
        color = _rgb(palette[index % len(palette)])
        chart_series.format.fill.solid()
        chart_series.format.fill.fore_color.rgb = color
        chart_series.format.line.color.rgb = color
    if normalized_chart_type in {"pie", "doughnut"} and chart.series:
        for index, point in enumerate(chart.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(palette[index % len(palette)])
            point.format.line.color.rgb = theme.white
    else:
        for axis in (chart.category_axis, chart.value_axis):
            _set_font_family(axis.tick_labels.font, theme.font_cn, theme.font_en)
            axis.tick_labels.font.size = Pt(float(label_size))
            axis.tick_labels.font.color.rgb = _rgb(resolved_label_color)
            axis.format.line.color.rgb = _rgb(resolved_grid_color)
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(resolved_grid_color)
    plot = chart.plots[0]
    plot.has_data_labels = bool(show_data_labels)
    if show_data_labels:
        _set_font_family(plot.data_labels.font, theme.font_cn, theme.font_en)
        plot.data_labels.font.size = Pt(float(label_size))
        plot.data_labels.font.color.rgb = _rgb(resolved_label_color)
    return chart


def _set_cell_border(
    cell,
    color: RGBColor | str | tuple[int, int, int],
    width: float,
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
        edge = etree.SubElement(properties, qn(tag))
        edge.set("w", str(int(Pt(float(width)))))
        solid_fill = etree.SubElement(edge, qn("a:solidFill"))
        rgb = etree.SubElement(solid_fill, qn("a:srgbClr"))
        rgb.set("val", str(_rgb(color)))


def _style_table_cell(
    cell,
    value: object,
    *,
    fill: RGBColor | str | tuple[int, int, int],
    color: RGBColor | str | tuple[int, int, int],
    size: float,
    bold: bool,
    align,
    margin: float,
    border_color: RGBColor | str | tuple[int, int, int],
    border_width: float,
    theme: Theme,
) -> None:
    cell.text = str(value)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    _set_fill(cell, fill)
    cell.margin_left = cell.margin_right = Emu(i_emu(margin))
    cell.margin_top = cell.margin_bottom = Emu(i_emu(margin))
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            _apply_run_font(
                run,
                size=size,
                bold=bold,
                color=color,
                font_cn=theme.font_cn,
                font_en=theme.font_en,
            )
    _set_cell_border(cell, border_color, border_width)


def editable_table(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    headers: Sequence[object],
    rows: Sequence[Sequence[object]],
    *,
    theme: Theme = DEFAULT_THEME,
    column_widths: Sequence[float] | None = None,
    header_size: float = 18,
    body_size: float = 16,
    header_fill: RGBColor | str | tuple[int, int, int] | None = None,
    header_color: RGBColor | str | tuple[int, int, int] | None = None,
    body_fill: RGBColor | str | tuple[int, int, int] | None = None,
    alternate_fill: RGBColor | str | tuple[int, int, int] | None = None,
    body_color: RGBColor | str | tuple[int, int, int] | None = None,
    border_color: RGBColor | str | tuple[int, int, int] | None = None,
    border_width: float = 0.5,
    cell_margin: float = 0.06,
    header_align=PP_ALIGN.CENTER,
    body_align=PP_ALIGN.LEFT,
):
    """Add a native, editable table with caller-supplied rows and headers."""
    headers = list(headers)
    rows = [list(row) for row in rows]
    if not headers:
        raise ValueError("editable_table requires at least one header")
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("each table row must match the header count")

    left, top, width, height = _box_emu(x, y, w, h)
    table = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height).table
    if column_widths is not None:
        widths = [_as_inches(value) for value in column_widths]
        if len(widths) != len(headers) or any(value <= 0 for value in widths):
            raise ValueError("column_widths must contain one positive width per column")
        if abs(sum(widths) - _as_inches(w)) > 1e-6:
            raise ValueError("column_widths must sum to the table width")
        width_emus = [i_emu(value) for value in widths[:-1]]
        width_emus.append(width - sum(width_emus))
        for column, column_width in zip(table.columns, width_emus):
            column.width = Emu(column_width)
    resolved_header_fill = theme.navy if header_fill is None else header_fill
    resolved_header_color = theme.white if header_color is None else header_color
    resolved_body_fill = theme.white if body_fill is None else body_fill
    resolved_alternate_fill = theme.paper if alternate_fill is None else alternate_fill
    resolved_body_color = theme.ink if body_color is None else body_color
    resolved_border_color = theme.line if border_color is None else border_color
    for column, value in enumerate(headers):
        _style_table_cell(
            table.cell(0, column),
            value,
            fill=resolved_header_fill,
            color=resolved_header_color,
            size=header_size,
            bold=True,
            align=header_align,
            margin=cell_margin,
            border_color=resolved_border_color,
            border_width=border_width,
            theme=theme,
        )
    for row_index, row in enumerate(rows, 1):
        fill = resolved_body_fill if row_index % 2 else resolved_alternate_fill
        for column, value in enumerate(row):
            _style_table_cell(
                table.cell(row_index, column),
                value,
                fill=fill,
                color=resolved_body_color,
                size=body_size,
                bold=False,
                align=body_align,
                margin=cell_margin,
                border_color=resolved_border_color,
                border_width=border_width,
                theme=theme,
            )
    return table


def timeline(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    milestones: Sequence[Mapping[str, object]],
    *,
    theme: Theme = DEFAULT_THEME,
    axis_color: RGBColor | str | tuple[int, int, int] | None = None,
    milestone_colors: Sequence[RGBColor | str | tuple[int, int, int]] | None = None,
    label_color: RGBColor | str | tuple[int, int, int] | None = None,
    title_color: RGBColor | str | tuple[int, int, int] | None = None,
    detail_color: RGBColor | str | tuple[int, int, int] | None = None,
    label_size: float = 12,
    title_size: float = 16,
    detail_size: float = 12,
    dot_diameter: float = 0.16,
    axis_weight: float = 1.5,
):
    """Draw milestones in the supplied order using fixed, even spacing."""
    items = list(milestones)
    if len(items) < 2:
        raise ValueError("timeline requires at least two milestones")
    for index, item in enumerate(items):
        if "label" not in item:
            raise ValueError(f"milestone {index} must contain an explicit label")
    x, y, w, h = (_as_inches(item) for item in (x, y, w, h))
    dot_diameter = _as_inches(dot_diameter)
    _box_emu(x, y, w, h)
    if w <= 0 or h <= 0 or dot_diameter <= 0:
        raise ValueError("timeline dimensions and dot_diameter must be positive")
    column_width = w / len(items)
    padding = 0.06
    if column_width <= padding * 2:
        raise ValueError("timeline width is too small for the milestone count")
    axis_y = y + h * 0.36
    centers = [
        x + column_width * (index + 0.5) for index in range(len(items))
    ]
    resolved_axis_color = theme.navy if axis_color is None else axis_color
    resolved_label_color = theme.accent if label_color is None else label_color
    resolved_title_color = theme.ink if title_color is None else title_color
    resolved_detail_color = theme.muted if detail_color is None else detail_color
    palette = _theme_palette(theme, milestone_colors)
    axis = line(slide, centers[0], axis_y, centers[-1], axis_y,
                color=resolved_axis_color, weight=axis_weight, theme=theme)
    dots = []
    labels = []
    for index, (item, center_x) in enumerate(zip(items, centers)):
        text_x = x + column_width * index + padding
        text_w = column_width - padding * 2
        labels.append(text_box(
            slide, text_x, y, text_w, h * 0.23, str(item["label"]),
            size=label_size, bold=True, color=resolved_label_color,
            font_cn=theme.font_cn, font_en=theme.font_en, theme=theme,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False,
        ))
        dots.append(
            circle(
                slide,
                center_x - dot_diameter / 2,
                axis_y - dot_diameter / 2,
                dot_diameter,
                fill=palette[index % len(palette)],
                line=theme.white,
                line_width=1,
                theme=theme,
            )
        )
        for top, height, key, size, color, bold, wrap in (
            (0.50, 0.18, "title", title_size, resolved_title_color, True, False),
            (0.70, 0.30, "detail", detail_size, resolved_detail_color, False, True),
        ):
            labels.append(text_box(
                slide, text_x, y + h * top, text_w, h * height,
                str(item.get(key, "")), size=size, bold=bold, color=color,
                font_cn=theme.font_cn, font_en=theme.font_en, theme=theme,
                align=PP_ALIGN.CENTER, wrap=wrap,
            ))
    return {"axis": axis, "dots": dots, "labels": labels}


def matrix(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    points: Sequence[Mapping[str, object]],
    *,
    x_label: str = "",
    y_label: str = "",
    theme: Theme = DEFAULT_THEME,
    background_color: RGBColor | str | tuple[int, int, int] | None = None,
    grid_color: RGBColor | str | tuple[int, int, int] | None = None,
    point_colors: Sequence[RGBColor | str | tuple[int, int, int]] | None = None,
    label_color: RGBColor | str | tuple[int, int, int] | None = None,
    axis_label_color: RGBColor | str | tuple[int, int, int] | None = None,
    label_size: float = 12,
    axis_label_size: float = 12,
    point_diameter: float = 0.14,
    point_label_width: float = 1.1,
):
    """Draw a normalized matrix from explicit ``label``, ``x``, and ``y`` values.

    ``x`` and ``y`` are inclusive normalized coordinates in the range 0..1;
    normalized ``y=1`` is the top of the plot.
    """

    points = list(points)
    if not points:
        raise ValueError("matrix requires at least one point")
    normalized = []
    for index, item in enumerate(points):
        missing = [key for key in ("label", "x", "y") if key not in item]
        if missing:
            raise ValueError(
                f"matrix point {index} is missing explicit fields: {', '.join(missing)}"
            )
        point_x = float(item["x"])
        point_y = float(item["y"])
        if not isfinite(point_x) or not isfinite(point_y):
            raise ValueError(f"matrix point {index} coordinates must be finite")
        if not 0.0 <= point_x <= 1.0 or not 0.0 <= point_y <= 1.0:
            raise ValueError(f"matrix point {index} coordinates must be within 0..1")
        normalized.append((str(item["label"]), point_x, point_y))
    x, y, w, h = (_as_inches(item) for item in (x, y, w, h))
    point_diameter = _as_inches(point_diameter)
    point_label_width = _as_inches(point_label_width)
    _box_emu(x, y, w, h)
    if w <= 0 or h <= 0 or point_diameter <= 0 or point_label_width <= 0:
        raise ValueError("matrix dimensions, point size, and label width must be positive")
    left_padding = 0.34 if y_label else 0.12
    right_padding = point_label_width + point_diameter / 2 + 0.08
    top_padding = 0.16
    bottom_padding = 0.34 if x_label else 0.16
    plot_x = x + left_padding
    plot_y = y + top_padding
    plot_w = w - left_padding - right_padding
    plot_h = h - top_padding - bottom_padding
    if plot_w <= 0 or plot_h <= 0:
        raise ValueError("matrix box is too small for its explicit labels and point size")
    resolved_background = theme.paper if background_color is None else background_color
    resolved_grid_color = theme.line if grid_color is None else grid_color
    resolved_label_color = theme.ink if label_color is None else label_color
    resolved_axis_label_color = (
        theme.muted if axis_label_color is None else axis_label_color
    )
    palette = _theme_palette(theme, point_colors)
    frame = rect(slide, x, y, w, h, fill=resolved_background,
                 line=resolved_grid_color, line_width=0.75, theme=theme)
    grid = [
        line(slide, plot_x + plot_w / 2, plot_y, plot_x + plot_w / 2,
             plot_y + plot_h, color=resolved_grid_color, weight=0.75, theme=theme),
        line(slide, plot_x, plot_y + plot_h / 2, plot_x + plot_w,
             plot_y + plot_h / 2, color=resolved_grid_color, weight=0.75, theme=theme),
        rect(slide, plot_x, plot_y, plot_w, plot_h, fill=None,
             line=resolved_grid_color, line_width=0.75, theme=theme),
    ]
    axis_labels = []
    if x_label:
        axis_labels.append(text_box(
            slide, plot_x, y + h - 0.28, plot_w, 0.22, x_label,
            size=axis_label_size, bold=True, color=resolved_axis_label_color,
            font_cn=theme.font_cn, font_en=theme.font_en, theme=theme,
            align=PP_ALIGN.CENTER, wrap=False,
        ))
    if y_label:
        y_axis_label = text_box(
            slide, x + 0.04, plot_y, 0.22, plot_h, y_label,
            size=axis_label_size, bold=True, color=resolved_axis_label_color,
            font_cn=theme.font_cn, font_en=theme.font_en, theme=theme,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False,
        )
        y_axis_label.rotation = 270
        axis_labels.append(y_axis_label)

    dots = []
    labels = []
    for index, (label, point_x, point_y) in enumerate(normalized):
        center_x = plot_x + plot_w * point_x
        center_y = plot_y + plot_h * (1.0 - point_y)
        dots.append(
            circle(
                slide,
                center_x - point_diameter / 2,
                center_y - point_diameter / 2,
                point_diameter,
                fill=palette[index % len(palette)],
                line=theme.white,
                line_width=0.8,
                theme=theme,
            )
        )
        labels.append(
            text_box(
                slide, center_x + point_diameter / 2 + 0.04, center_y - 0.12,
                point_label_width, 0.24, label, size=label_size, bold=True,
                color=resolved_label_color, font_cn=theme.font_cn,
                font_en=theme.font_en, theme=theme,
                anchor=MSO_ANCHOR.MIDDLE, wrap=False,
            )
        )
    return {
        "frame": frame,
        "grid": grid,
        "axis_labels": axis_labels,
        "dots": dots,
        "labels": labels,
    }


def out_of_bounds(slide, *, width: int = SLIDE_W, height: int = SLIDE_H) -> list[str]:
    """Return shape names whose geometry exceeds the slide canvas."""

    errors: list[str] = []
    for index, shape in enumerate(slide.shapes, 1):
        left = int(getattr(shape, "left", 0) or 0)
        top = int(getattr(shape, "top", 0) or 0)
        shape_width = int(getattr(shape, "width", 0) or 0)
        shape_height = int(getattr(shape, "height", 0) or 0)
        if (
            left < 0
            or top < 0
            or shape_width < 0
            or shape_height < 0
            or left + shape_width > int(width)
            or top + shape_height > int(height)
        ):
            name = getattr(shape, "name", "") or f"shape-{index}"
            errors.append(name)
    return errors


def assert_slide_bounds(slide, *, label: str = "slide") -> None:
    errors = out_of_bounds(slide)
    if errors:
        raise ValueError(f"{label} has out-of-bounds shapes: {', '.join(errors)}")


def looks_like_serialized_rich_text_payload(text: str) -> bool:
    compact = " ".join(str(text).split())
    dict_payload = (
        compact.startswith(("[{'text':", '[{"text":'))
        and ("'size':" in compact or '"size":' in compact)
        and ("'color':" in compact or '"color":' in compact)
        and ("'bold':" in compact or '"bold":' in compact)
    )
    enum_payload = (
        "<PP_PARAGRAPH_ALIGNMENT." in compact
        or "<MSO_VERTICAL_ANCHOR." in compact
    )
    return dict_payload or enum_payload


def assert_deck_text_sanity(deck: Presentation) -> None:
    failures: list[str] = []
    for slide_index, slide in enumerate(deck.slides, 1):
        for shape_index, shape in enumerate(slide.shapes, 1):
            if not getattr(shape, "has_text_frame", False):
                continue
            text = getattr(shape, "text", "") or ""
            if looks_like_serialized_rich_text_payload(text):
                failures.append(f"slide {slide_index} shape {shape_index}")
    if failures:
        raise ValueError(
            "serialized rich-text/style payload leaked into audience text: "
            + ", ".join(failures)
        )


def assert_deck_bounds(deck: Presentation) -> None:
    for index, slide in enumerate(deck.slides, 1):
        assert_slide_bounds(slide, label=f"slide {index}")
    assert_deck_text_sanity(deck)


__all__ = [
    "DEFAULT_THEME",
    "GOOGLE_CLEAN_THEME",
    "DEFAULT_SEMANTIC_LEGEND",
    "EMU_PER_INCH",
    "Edge",
    "LayoutRegion",
    "SLIDE_H",
    "SLIDE_H_IN",
    "SLIDE_W",
    "SLIDE_W_IN",
    "Theme",
    "add_background",
    "assert_deck_bounds",
    "assert_deck_text_sanity",
    "assert_slide_bounds",
    "card",
    "chip",
    "circle",
    "decision_frame",
    "editable_table",
    "edges_first",
    "header_footer",
    "i_emu",
    "line",
    "layout_regions",
    "matrix",
    "native_chart",
    "new_presentation",
    "new_slide",
    "out_of_bounds",
    "proportional_bar",
    "rect",
    "rich_text_box",
    "round_rect",
    "semantic_legend",
    "text_box",
    "timeline",
]
