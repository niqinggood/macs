#!/usr/bin/env python3
"""Render an editable task deck from a content contract and page compositions.

The JSON file owns facts, narrative jobs, and evidence.  In a scaffolded
audience-facing build, the task-local builder must also be edited so every
slide is assigned an explicit renderer.  The stable renderers below remain
reusable mechanics; they are not the whole-deck layout engine.
"""

from __future__ import annotations

from collections import Counter
import json
from pathlib import Path
import unicodedata

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from pptx_design import (
    DEFAULT_THEME,
    GOOGLE_CLEAN_THEME,
    Edge,
    Theme,
    assert_deck_bounds,
    circle,
    decision_frame,
    editable_table,
    edges_first,
    header_footer,
    image_contain,
    image_cover,
    line,
    layout_regions,
    matrix,
    native_chart,
    new_presentation,
    new_slide,
    proportional_bar,
    rect,
    rich_text_box,
    round_rect,
    semantic_legend,
    text_box,
    timeline as timeline_diagram,
)


OUTPUT_PPTX = Path("__TASK_OUTPUT_PPTX__")
COMPOSITION_MODE = "__TASK_COMPOSITION_MODE__"
SPEC_PATH = Path(__file__).with_name("deck_spec.json")
VISUAL_PLAN_PATH = Path(__file__).with_name("visual_plan.json")
THEME = DEFAULT_THEME
VISUAL_PLAN = {}
VISUAL_DIRECTION = {}
SLIDE_VISUALS = {}


def planned_layout(page: int):
    """Turn the selected per-slide layout into real, editable content geometry."""
    visual = SLIDE_VISUALS.get(page, {})
    family = visual.get("layout_family")
    if not family:
        raise ValueError(f"slide {page} needs an explicit layout_family before planned_layout()")
    return layout_regions(
        family,
        primary_share=visual.get("primary_share", 0.68),
        mirror=visual.get("mirror", False),
    )


ALLOWED_KINDS = {
    "cover", "statement", "cards", "split", "process", "metrics",
    "chart", "table", "matrix", "timeline", "closing",
}
EXPLANATORY_KINDS = {"process", "metrics", "chart", "table", "matrix", "timeline"}


def _text(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (str, int, float)):
        return str(value).strip()
    # Containers are structured production data, never audience copy.  Silent
    # ``str(dict)`` conversion previously leaked brand/action dictionaries onto
    # rendered slides while still producing a technically valid PPTX.
    return ""


def _items(value) -> list:
    return list(value) if isinstance(value, list) else []


def _bullets(values) -> str:
    return "\n".join(f"• {_text(item)}" for item in _items(values) if _text(item))


def _brand_name(value) -> str:
    if isinstance(value, dict):
        for key in ("name", "brand", "title", "label"):
            text = _text(value.get(key))
            if text:
                return text
        return ""
    return _text(value)


def _closing_action(value, index: int) -> dict[str, str]:
    if isinstance(value, dict):
        title = (
            _text(value.get("title"))
            or _text(value.get("owner"))
            or _text(value.get("label"))
            or f"行动 {index}"
        )
        body = (
            _text(value.get("body"))
            or _text(value.get("action"))
            or _text(value.get("desc"))
            or _text(value.get("detail"))
        )
        return {"title": title, "body": body}
    return {"title": f"行动 {index}", "body": _text(value)}


def _closing_actions(slide: dict) -> list[dict[str, str]]:
    values = _items(slide.get("items")) or _items(slide.get("next_actions"))
    return [_closing_action(value, index) for index, value in enumerate(values[:3], 1)]


def _closing_commitments(slide: dict) -> list[dict[str, str]]:
    return [
        _closing_action(value, index)
        for index, value in enumerate(_items(slide.get("commitments"))[:3], 1)
    ]


def _cards_conclusion(slide: dict) -> tuple[str, str]:
    value = slide.get("conclusion")
    if isinstance(value, dict):
        label = _text(value.get("title")) or _text(value.get("label")) or "结论"
        body = (
            _text(value.get("body"))
            or _text(value.get("text"))
            or _text(value.get("claim"))
            or _text(value.get("takeaway"))
        )
        if body:
            return label, body
    else:
        body = _text(value)
        if body:
            return "结论", body
    return "结论", _text(slide.get("takeaway"))


def _timeline_milestone(value, index: int) -> dict[str, str]:
    if not isinstance(value, dict):
        return {"label": str(index), "title": _text(value), "detail": ""}
    label = _text(value.get("label")) or _text(value.get("phase")) or _text(value.get("time"))
    title = _text(value.get("title"))
    if not title and "·" in label:
        label, title = (part.strip() for part in label.split("·", 1))
    milestone = _text(value.get("milestone"))
    points = [_text(point) for point in _items(value.get("points")) if _text(point)]
    detail = _text(value.get("detail")) or _text(value.get("body"))
    if milestone:
        detail = milestone + (("\n" + "\n".join(f"• {point}" for point in points)) if points else "")
    elif points:
        detail = "\n".join(f"• {point}" for point in points)
    return {
        "label": label or str(index),
        "title": title or milestone or f"阶段 {index}",
        "detail": detail,
    }


def _statement_claim(slide: dict) -> str:
    """Accept the concise aliases models naturally use for a statement page."""

    return (
        _text(slide.get("claim"))
        or _text(slide.get("takeaway"))
        or _text(slide.get("body"))
        or _text(slide.get("statement"))
    )


def _statement_evidence(slide: dict) -> list:
    return _items(slide.get("evidence")) or _items(slide.get("points"))


def _title_size(value: str) -> int:
    length = len(_text(value))
    if length <= 22:
        return 40
    if length <= 34:
        return 37
    return 35


def _body_size(value: str, *, roomy: int = 18, compact: int = 16) -> int:
    length = len(_text(value))
    if length <= 52:
        return roomy
    if length <= 90:
        return compact
    return compact


def _visual_units(value: str) -> float:
    """Estimate rendered headline width across CJK and Latin scripts."""

    units = 0.0
    for character in _text(value):
        if character.isspace():
            units += 0.32
        elif unicodedata.east_asian_width(character) in {"W", "F"}:
            units += 1.0
        elif character.isalnum():
            units += 0.58
        else:
            units += 0.68
    return max(1.0, units)


def _headline_size(
    value: str,
    width: float,
    *,
    maximum: int = 42,
    minimum: int = 35,
) -> int:
    """Keep conclusion headlines large without orphaning a final glyph.

    Character count treats a Chinese glyph and a Latin letter as equal even
    though their rendered widths differ substantially.  Fit concise headlines
    to one line when the formal 35pt floor permits it; longer headlines retain
    a restrained two-line composition instead of overflowing the slide.
    """

    units = _visual_units(value)
    usable_points = width * 72 * 0.96
    if units <= 30:
        fitted = int(usable_points / units)
    else:
        fitted = int((usable_points * 1.72) / units)
    return max(minimum, min(maximum, fitted))


def _metric_value_style(value: str, width: float) -> tuple[str, int]:
    """Keep metric values legible without colliding with their separator.

    Mixed Chinese/currency values are wider than a raw character count suggests.
    Estimate their visual width, shrink only the headline value when necessary,
    and preserve the 14pt+ body hierarchy.
    """

    compact = _text(value).replace(" 万", "万").replace(" 亿", "亿").replace(" %", "%")
    visual_units = 0.0
    for character in compact:
        if "\u2e80" <= character <= "\u9fff":
            visual_units += 1.0
        elif character.isalnum():
            visual_units += 0.58
        elif character.isspace():
            visual_units += 0.3
        else:
            visual_units += 0.68
    available_points = max(1.0, (width - 0.28) * 72 * 0.82)
    adaptive = int(available_points / max(1.0, visual_units))
    return compact, max(22, min(38, adaptive))


def _validate_slide_content(slide: dict, index: int, kind: str) -> None:
    """Reject only content volumes that cannot stay readable at 14pt+."""

    def limited(value, label: str, maximum: int) -> None:
        if len(_text(value)) > maximum:
            raise ValueError(
                f"slide {index} {label} exceeds {maximum} characters; split or prioritize it"
            )

    if kind == "statement":
        claim = _statement_claim(slide)
        if not claim:
            raise ValueError(
                f"slide {index} statement requires claim, takeaway, body, or statement"
            )
        limited(claim, "claim", 120)
        evidence = _statement_evidence(slide)
        if len(evidence) > 4:
            raise ValueError(f"slide {index} statement supports at most 4 evidence items")
        for offset, value in enumerate(evidence, 1):
            limited(value, f"evidence[{offset}]", 80)
    elif kind == "cards":
        items = _items(slide.get("items"))
        if not 2 <= len(items) <= 6:
            raise ValueError(f"slide {index} cards requires 2 to 6 items")
        conclusion_label, conclusion_body = _cards_conclusion(slide)
        limited(conclusion_label, "conclusion.title", 12)
        limited(conclusion_body, "conclusion.body", 90)
        body_limit = 90 if len(items) <= 3 else (50 if len(items) == 4 else 38)
        for offset, item in enumerate(items, 1):
            if isinstance(item, dict):
                limited(item.get("title") or item.get("label"), f"items[{offset}].title", 42)
                limited(item.get("body") or item.get("desc") or item.get("note"), f"items[{offset}].body", body_limit)
    elif kind == "split":
        for side, maximum in (("left", 130), ("right", 90)):
            panel = slide.get(side) or {}
            if isinstance(panel, dict):
                limited(panel.get("title"), f"{side}.title", 48)
                body = "\n".join(
                    part for part in (_text(panel.get("body")), _bullets(panel.get("bullets"))) if part
                )
                limited(body, f"{side}.body", maximum)
    elif kind == "process":
        items = _items(slide.get("items"))
        if not 3 <= len(items) <= 5:
            raise ValueError(f"slide {index} process requires 3 to 5 items")
        for offset, item in enumerate(items, 1):
            if isinstance(item, dict):
                limited(item.get("title"), f"items[{offset}].title", 30)
                limited(item.get("body") or item.get("desc") or item.get("detail"), f"items[{offset}].body", 48)
        lateral = slide.get("lateral")
        if lateral is not None and not isinstance(lateral, list):
            raise ValueError(f"slide {index} process lateral must be a list")
        lateral_items = _items(lateral)
        if len(lateral_items) > 4:
            raise ValueError(f"slide {index} process supports at most 4 lateral capabilities")
        for offset, value in enumerate(lateral_items, 1):
            limited(value, f"lateral[{offset}]", 24)
    elif kind == "metrics":
        items = _items(slide.get("items")) or _items(slide.get("metrics"))
        if not 2 <= len(items) <= 4:
            raise ValueError(f"slide {index} metrics requires 2 to 4 items")
        for offset, item in enumerate(items, 1):
            if isinstance(item, dict):
                limited(item.get("value"), f"items[{offset}].value", 26)
                limited(item.get("label"), f"items[{offset}].label", 36)
                limited(item.get("body") or item.get("note") or item.get("desc"), f"items[{offset}].body", 72)
    elif kind == "chart":
        categories = _items(slide.get("categories"))
        series = _items(slide.get("series"))
        if not 2 <= len(categories) <= 8 or not 1 <= len(series) <= 4:
            raise ValueError(f"slide {index} chart needs 2 to 8 categories and 1 to 4 series")
        for offset, item in enumerate(series, 1):
            if not isinstance(item, dict) or len(_items(item.get("values"))) != len(categories):
                raise ValueError(f"slide {index} series[{offset}] values must match categories")
        # Notes sit in a dedicated editorial callout and are independently
        # checked by the rendered-overflow auditor. A 150-character hard stop
        # rejected otherwise valid Chinese funding breakdowns by one or two
        # characters and caused needless model rewrite loops.
        limited(slide.get("note") or slide.get("intro"), "note", 180)
    elif kind == "table":
        headers = _items(slide.get("headers"))
        rows = _items(slide.get("rows"))
        if not 2 <= len(headers) <= 7 or not 1 <= len(rows) <= 6:
            raise ValueError(f"slide {index} table needs 2 to 7 columns and 1 to 6 rows")
        for row_offset, row in enumerate(rows, 1):
            if not isinstance(row, list) or len(row) != len(headers):
                raise ValueError(f"slide {index} row {row_offset} must match the header count")
            for col_offset, value in enumerate(row, 1):
                limited(value, f"rows[{row_offset}][{col_offset}]", 70)
    elif kind == "matrix":
        points = _items(slide.get("points"))
        if points:
            if len(points) > 8:
                raise ValueError(f"slide {index} matrix supports at most 8 points")
            for offset, point in enumerate(points, 1):
                if isinstance(point, dict):
                    limited(point.get("label"), f"points[{offset}].label", 42)
    elif kind == "timeline":
        items = _items(slide.get("items")) or _items(slide.get("milestones"))
        if not 2 <= len(items) <= 6:
            raise ValueError(f"slide {index} timeline requires 2 to 6 milestones")
        for offset, item in enumerate(items, 1):
            if isinstance(item, dict):
                normalized = _timeline_milestone(item, offset)
                limited(normalized.get("label"), f"items[{offset}].label", 28)
                limited(normalized.get("title"), f"items[{offset}].title", 34)
                limited(normalized.get("detail"), f"items[{offset}].detail", 120)
                points = _items(item.get("points"))
                if len(points) > 4:
                    raise ValueError(f"slide {index} items[{offset}].points supports at most 4 entries")
                for point_offset, point in enumerate(points, 1):
                    limited(point, f"items[{offset}].points[{point_offset}]", 28)
    elif kind == "closing":
        items = _items(slide.get("items")) or _items(slide.get("next_actions"))
        if len(items) > 3:
            raise ValueError(f"slide {index} closing supports at most 3 actions")
        for offset, item in enumerate(items, 1):
            action = _closing_action(item, offset)
            limited(action.get("title"), f"actions[{offset}].title", 42)
            limited(action.get("body"), f"actions[{offset}].body", 72)
        commitments = _items(slide.get("commitments"))
        if len(commitments) > 3:
            raise ValueError(f"slide {index} closing supports at most 3 commitments")
        limited(slide.get("decisions_label"), "decisions_label", 48)
        limited(slide.get("commitments_label"), "commitments_label", 48)
        for offset, item in enumerate(commitments, 1):
            commitment = _closing_action(item, offset)
            limited(commitment.get("title"), f"commitments[{offset}].title", 24)
            limited(commitment.get("body"), f"commitments[{offset}].body", 62)


def _validate_evidence_contract(spec: dict) -> None:
    """Require explicit slide-to-source traceability in the formal builder.

    The host independently verifies excerpts against trusted tool evidence.
    This local check keeps the saved builder honest even when it is run outside
    the agent loop later.
    """

    art_direction = spec.get("art_direction")
    if not isinstance(art_direction, dict):
        raise ValueError("formal per-slide builds require an art_direction object")
    required_direction = (
        "visual_thesis", "composition_strategy", "asset_strategy",
        "palette_strategy", "typography_strategy", "background_rhythm",
    )
    missing_direction = [key for key in required_direction if not _text(art_direction.get(key))]
    if missing_direction:
        raise ValueError(
            "formal art_direction is incomplete; missing: " + ", ".join(missing_direction)
        )

    policy = _text(spec.get("evidence_policy")).lower()
    if policy not in {"source-bound", "evidence-bound", "事实约束", "证据约束"}:
        raise ValueError("formal per-slide builds require evidence_policy=source-bound")
    entries = spec.get("evidence_ledger")
    if not isinstance(entries, list):
        raise ValueError("source-bound deck_spec.json requires an evidence_ledger array")
    evidence_ids: set[str] = set()
    for index, entry in enumerate(entries, 1):
        if not isinstance(entry, dict):
            raise ValueError(f"evidence_ledger[{index}] must be an object")
        evidence_id = _text(entry.get("id"))
        if not evidence_id or evidence_id in evidence_ids:
            raise ValueError(f"evidence_ledger[{index}] needs a unique id")
        for field in ("source", "locator", "excerpt"):
            if not _text(entry.get(field)):
                raise ValueError(f"evidence_ledger[{index}] is missing {field}")
        evidence_ids.add(evidence_id)

    source_markers = (
        "source: user", "source: local", "source: project", "source: public",
        "来源：用户", "来源：本地", "来源：项目", "公开来源", "用户陈述", "用户材料", "本地材料",
    )
    non_factual_markers = (
        "规划估算（非项目事实）", "planning estimate (not a project fact)",
        "待确认", "未统计", "未提供", "未知", "不含新增事实", "no new factual claims",
    )
    for index, slide in enumerate(_items(spec.get("slides")), 1):
        if index == 1 or _text(slide.get("kind")).lower() in {"cover", "title"}:
            continue
        note = _text(slide.get("source_note"))
        lower_note = note.lower()
        if not note:
            raise ValueError(f"slide {index} requires a visible source_note")
        source_backed = any(marker in lower_note for marker in source_markers)
        non_factual = any(marker.lower() in lower_note for marker in non_factual_markers)
        if source_backed and not non_factual:
            refs = slide.get("evidence_refs")
            if not isinstance(refs, list) or not refs:
                raise ValueError(f"slide {index} requires evidence_refs")
            missing = [_text(ref) for ref in refs if _text(ref) not in evidence_ids]
            if missing:
                raise ValueError(f"slide {index} has unresolved evidence_refs: {missing}")


def _theme_for_spec(spec: dict, direction: dict | None = None):
    direction = direction if isinstance(direction, dict) else {}
    palette = direction.get("palette")
    typography = direction.get("typography")
    if isinstance(palette, dict) and isinstance(typography, dict):
        def color(key: str, fallback: RGBColor) -> RGBColor:
            raw = _text(palette.get(key)).lstrip("#")
            if len(raw) != 6:
                return fallback
            try:
                return RGBColor.from_string(raw.upper())
            except ValueError:
                return fallback

        return Theme(
            navy=color("navy", DEFAULT_THEME.navy),
            ink=color("ink", DEFAULT_THEME.ink),
            muted=color("muted", DEFAULT_THEME.muted),
            line=color("line", DEFAULT_THEME.line),
            paper=color("paper", DEFAULT_THEME.paper),
            white=color("white", DEFAULT_THEME.white),
            accent=color("accent", DEFAULT_THEME.accent),
            accent_soft=color("accent_soft", DEFAULT_THEME.accent_soft),
            positive=color("positive", DEFAULT_THEME.positive),
            positive_soft=color("positive_soft", DEFAULT_THEME.positive_soft),
            warning=color("warning", DEFAULT_THEME.warning),
            warning_soft=color("warning_soft", DEFAULT_THEME.warning_soft),
            negative=color("negative", DEFAULT_THEME.negative),
            negative_soft=color("negative_soft", DEFAULT_THEME.negative_soft),
            font_cn=_text(typography.get("font_cn")) or DEFAULT_THEME.font_cn,
            font_en=_text(typography.get("font_en")) or DEFAULT_THEME.font_en,
        )
    requested = _text(spec.get("theme") or spec.get("style")).lower().replace("-", "_")
    if requested in {"google", "google_clean", "bright", "light_product"}:
        return GOOGLE_CLEAN_THEME
    return DEFAULT_THEME


def _load_visual_plan(total: int) -> dict:
    if COMPOSITION_MODE != "per-slide":
        return {}
    if not VISUAL_PLAN_PATH.is_file():
        raise ValueError("formal per-slide builds require task-local visual_plan.json")
    plan = json.loads(VISUAL_PLAN_PATH.read_text(encoding="utf-8"))
    if plan.get("needs_customization", True):
        raise ValueError("visual_plan.json still needs task-specific direction and assets")
    serialized = json.dumps(plan, ensure_ascii=False).lower()
    if "replace with" in serialized or "待替换" in serialized:
        raise ValueError("visual_plan.json still contains template placeholder text")
    direction = plan.get("selected_direction")
    if not isinstance(direction, dict):
        raise ValueError("visual_plan.json requires selected_direction")
    for field in ("id", "name", "visual_thesis", "rationale", "image_style", "chart_style"):
        if not _text(direction.get(field)):
            raise ValueError(f"visual_plan selected_direction is missing {field}")
    typography = direction.get("typography")
    if not isinstance(typography, dict):
        raise ValueError("visual_plan selected_direction requires typography")
    for field, floor in (("deck_title_pt", 50), ("slide_title_pt", 35), ("subhead_pt", 24), ("body_pt", 16)):
        value = typography.get(field)
        if not isinstance(value, (int, float)) or isinstance(value, bool) or value < floor:
            raise ValueError(f"visual_plan typography.{field} must be at least {floor}")
    palette = direction.get("palette")
    if not isinstance(palette, dict):
        raise ValueError("visual_plan selected_direction requires palette")
    for field in (
        "navy", "ink", "muted", "line", "paper", "white", "accent", "accent_soft",
        "positive", "positive_soft", "warning", "warning_soft", "negative", "negative_soft",
    ):
        value = _text(palette.get(field))
        if len(value) != 7 or not value.startswith("#"):
            raise ValueError(f"visual_plan palette.{field} must be #RRGGBB")
        try:
            int(value[1:], 16)
        except ValueError as exc:
            raise ValueError(f"visual_plan palette.{field} must be #RRGGBB") from exc
    slides = _items(plan.get("slides"))
    if len(slides) != total:
        raise ValueError(f"visual_plan has {len(slides)} slides; deck_spec has {total}")
    seen = set()
    for index, item in enumerate(slides, 1):
        if not isinstance(item, dict) or item.get("slide") != index:
            raise ValueError(f"visual_plan slides[{index}] must describe slide {index}")
        for field in ("narrative_job", "composition", "density", "focal_anchor"):
            if not _text(item.get(field)):
                raise ValueError(f"visual_plan slides[{index}] is missing {field}")
        asset = item.get("asset")
        if not isinstance(asset, dict) or not _text(asset.get("type")):
            raise ValueError(f"visual_plan slides[{index}] requires one primary asset")
        for field in ("purpose", "source"):
            if not _text(asset.get(field)):
                raise ValueError(f"visual_plan slides[{index}].asset is missing {field}")
        if _text(asset.get("type")).lower() in {"photo", "illustration", "screenshot"}:
            raw_path = _text(asset.get("path"))
            source = Path(raw_path).expanduser()
            if not source.is_absolute():
                source = VISUAL_PLAN_PATH.parent / source
            if not raw_path or not source.is_file():
                raise ValueError(f"visual_plan slide {index} media asset does not exist: {raw_path}")
        seen.add(index)
    if seen != set(range(1, total + 1)):
        raise ValueError("visual_plan slide coverage is incomplete")
    return plan


def planned_visual(page: int) -> dict:
    """Return one slide's validated visual-plan entry to a task renderer."""

    return dict(SLIDE_VISUALS.get(int(page), {}))


def planned_image(slide, page: int, x: float, y: float, w: float, h: float, *, mode: str | None = None):
    """Place the media asset assigned to ``page`` without stretching it."""

    visual = SLIDE_VISUALS.get(int(page), {})
    asset = visual.get("asset") if isinstance(visual, dict) else None
    if not isinstance(asset, dict) or _text(asset.get("type")).lower() not in {"photo", "illustration", "screenshot"}:
        raise ValueError(f"visual_plan slide {page} does not assign a media asset")
    raw_path = _text(asset.get("path"))
    source = Path(raw_path).expanduser()
    if not source.is_absolute():
        source = VISUAL_PLAN_PATH.parent / source
    crop = (mode or _text(asset.get("crop")) or "contain").lower()
    if crop == "contain":
        return image_contain(slide, source, x, y, w, h)
    return image_cover(
        slide, source, x, y, w, h,
        focal_x=float(asset.get("focal_x", 0.5)),
        focal_y=float(asset.get("focal_y", 0.5)),
    )


def load_spec() -> dict:
    spec = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
    if spec.get("needs_customization", True):
        raise ValueError("deck_spec.json still needs task content")
    slides = _items(spec.get("slides"))
    if not 6 <= len(slides) <= 14:
        raise ValueError("deck_spec.json must contain 6 to 14 slides")
    if _text(slides[0].get("kind")).lower() != "cover":
        raise ValueError("the first slide must use kind=cover")
    quality_profile = _text(spec.get("quality_profile") or "audience-facing").lower()
    if quality_profile not in {"audience-facing", "recurring"}:
        raise ValueError("quality_profile must be audience-facing or recurring")
    if quality_profile == "audience-facing" and _text(slides[-1].get("kind")).lower() != "closing":
        raise ValueError("an audience-facing deck must end with kind=closing")
    body_kinds = []
    for index, slide in enumerate(slides, 1):
        kind = _text(slide.get("kind")).lower()
        if kind not in ALLOWED_KINDS:
            raise ValueError(f"slide {index} has unsupported kind: {kind!r}")
        if not _text(slide.get("title")):
            raise ValueError(f"slide {index} is missing a title")
        if index not in {1, len(slides)}:
            body_kinds.append(kind)
        for key in ("body", "takeaway", "claim", "note", "intro"):
            value = slide.get(key)
            if isinstance(value, str) and len(value.strip()) > 220:
                raise ValueError(f"slide {index} field {key!r} is too dense; split or prioritize it")
        _validate_slide_content(slide, index, kind)
    # The formal per-slide builder carries the source-traceability contract.
    # The stable builder is also reused by the lightweight one-command helper,
    # whose specs may be original-content-only and therefore have no sources.
    if COMPOSITION_MODE == "per-slide":
        _validate_evidence_contract(spec)
    if quality_profile == "audience-facing" and len(slides) >= 8:
        counts = Counter(body_kinds)
        minimum_families = 5 if len(slides) >= 10 else 4
        if len(counts) < minimum_families:
            raise ValueError(f"audience-facing deck needs {minimum_families} body composition kinds")
        if counts.get("cards", 0) > 1:
            raise ValueError("cards may appear at most once in an audience-facing deck")
        if not any(kind in EXPLANATORY_KINDS for kind in body_kinds):
            raise ValueError("audience-facing deck needs an explanatory visual composition")
        if any(left == middle == right for left, middle, right in zip(body_kinds, body_kinds[1:], body_kinds[2:])):
            raise ValueError("three consecutive content slides must not use the same composition kind")
    return spec


def _section_slide(deck, spec: dict, slide_spec: dict, page: int, total: int):
    slide = new_slide(deck)
    title = _text(slide_spec.get("title"))
    header_footer(
        slide,
        brand=_brand_name(spec.get("brand")) or "OpenKitty",
        title=title,
        section=_text(slide_spec.get("section")) or _text(slide_spec.get("kind")).upper(),
        page=page,
        total=total,
        theme=THEME,
    )
    takeaway = _text(slide_spec.get("takeaway"))
    if takeaway:
        rect(slide, 0.58, 1.40, 0.075, 0.52, fill=THEME.accent)
        text_box(slide, 0.84, 1.39, 11.8, 0.54, takeaway, size=17, color=THEME.muted, anchor=MSO_ANCHOR.MIDDLE)
    source_note = _text(slide_spec.get("source_note"))
    if source_note:
        text_box(slide, 0.58, 6.88, 11.8, 0.18, source_note, size=8, color=THEME.muted, wrap=False)
    return slide


def render_cover(deck, spec: dict, page_spec: dict) -> None:
    slide = new_slide(deck, THEME.navy)
    rect(slide, 0, 2.02, 0.12, 3.38, fill=THEME.accent)
    rect(slide, 0, 5.40, 13.333, 0.075, fill=THEME.accent)
    rect(slide, 0.72, 6.43, 0.92, 0.045, fill=THEME.accent)
    kicker = _text(page_spec.get("kicker")) or "EXECUTIVE BRIEF"
    kicker_size = 11 if len(kicker) <= 28 else 9 if len(kicker) <= 58 else 8
    kicker_height = 0.32 if len(kicker) <= 28 else 0.78
    text_box(slide, 0.72, 0.58, 5.2, kicker_height, kicker, size=kicker_size, bold=True, color=THEME.accent, wrap=True)
    title = _text(page_spec.get("title"))
    cover_size = 58 if len(title) <= 16 else 54 if len(title) <= 26 else 50
    text_box(slide, 0.72, 2.05, 11.9, 1.82, title, size=cover_size, bold=True, color=THEME.white, anchor=MSO_ANCHOR.MIDDLE)
    subtitle = _text(page_spec.get("subtitle")) or _text(spec.get("subtitle"))
    text_box(slide, 0.72, 4.28, 11.4, 0.72, subtitle, size=19, color=THEME.accent)
    text_box(slide, 0.72, 6.72, 4.8, 0.34, _brand_name(spec.get("brand")), size=16, bold=True, color=THEME.white, wrap=False)
    meta = _text(page_spec.get("meta")) or _text(spec.get("meta"))
    if meta:
        text_box(slide, 7.0, 6.72, 5.55, 0.34, meta, size=16, color=THEME.accent, align=PP_ALIGN.RIGHT, wrap=False)


def render_statement(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    claim = _statement_claim(page_spec)
    rect(slide, 0.66, 2.28, 0.11, 1.7, fill=THEME.accent)
    claim_size = 31 if len(claim) <= 45 else 27
    text_box(slide, 1.02, 2.25, 10.95, 1.72, claim, size=claim_size, bold=True, color=THEME.navy, anchor=MSO_ANCHOR.MIDDLE)
    evidence = _statement_evidence(page_spec)[:4]
    if evidence:
        width = 11.55 / len(evidence)
        start = 0.85
        for index, item in enumerate(evidence):
            x = start + index * width
            if index:
                line(slide, x, 4.55, x, 5.95, color=THEME.line, weight=1)
            text_box(slide, x + 0.14, 4.5, width - 0.28, 0.28, f"0{index + 1}", size=11, bold=True, color=THEME.accent, wrap=False)
            text_box(slide, x + 0.14, 5.02, width - 0.28, 0.86, _text(item), size=17, color=THEME.ink, align=PP_ALIGN.LEFT)


def render_cards(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    conclusion_label, conclusion = _cards_conclusion(page_spec)
    section_spec = dict(page_spec)
    if conclusion:
        # Cards use the takeaway as a high-contrast conclusion rail instead of
        # a low-emphasis subtitle. This prevents a required conclusion from
        # looking like optional metadata or disappearing in a dense grid.
        section_spec["takeaway"] = ""
    slide = _section_slide(deck, spec, section_spec, page, total)
    items = _items(page_spec.get("items"))[:6]
    if not items:
        raise ValueError(f"slide {page} cards requires items")
    # Editorial columns for small sets; two-column evidence bands for larger
    # sets.  Both avoid the repeated rounded-card dashboard silhouette.
    columns = len(items) if len(items) <= 3 else (2 if len(items) == 4 else 3)
    rows = 1 if len(items) <= 3 else 2
    gap_x = 0.38
    width = (12.15 - gap_x * (columns - 1)) / columns
    available_height = 3.63 if conclusion else 3.95
    row_gap = 0.18 if conclusion else 0.2
    height = available_height if rows == 1 else (available_height - row_gap * (rows - 1)) / rows
    for index, item in enumerate(items):
        row, column = divmod(index, columns)
        x = 0.58 + column * (width + gap_x)
        y = 2.35 + row * (height + row_gap)
        if rows > 1:
            rect(slide, x, y, width, height, fill=THEME.paper if (row + column) % 2 == 0 else THEME.white)
        elif index:
            line(slide, x - gap_x / 2, y + 0.15, x - gap_x / 2, y + height - 0.1, color=THEME.line, weight=1)
        metric = (_text(item.get("metric")) or _text(item.get("value"))) if isinstance(item, dict) else ""
        title = (_text(item.get("title")) or _text(item.get("label"))) if isinstance(item, dict) else _text(item)
        body = (_text(item.get("body")) or _text(item.get("desc")) or _text(item.get("note"))) if isinstance(item, dict) else ""
        text_box(slide, x + 0.22, y + 0.16, 0.52, 0.26, f"{index + 1:02d}", size=11, bold=True, color=THEME.accent, wrap=False)
        if metric:
            text_box(slide, x + 0.82, y + 0.11, width - 1.05, 0.4, metric, size=23, bold=True, color=THEME.accent)
            title_y = y + 0.62
        else:
            title_y = y + 0.48
        text_box(slide, x + 0.22, title_y, width - 0.44, 0.42, title, size=22 if rows == 1 else 18, bold=True, color=THEME.navy)
        body_y = title_y + 0.58
        text_box(slide, x + 0.22, body_y, width - 0.44, max(0.42, y + height - body_y - 0.14), body, size=_body_size(body), color=THEME.muted)
    if conclusion:
        rect(slide, 0.58, 6.10, 12.15, 0.58, fill=THEME.navy)
        text_box(slide, 0.80, 6.19, 0.72, 0.38, conclusion_label, size=14, bold=True, color=THEME.accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        line(slide, 1.68, 6.19, 1.68, 6.58, color=THEME.accent, weight=1)
        text_box(slide, 1.92, 6.17, 10.45, 0.40, conclusion, size=17, bold=True, color=THEME.white, anchor=MSO_ANCHOR.MIDDLE)


def render_split(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    halves = [page_spec.get("left") or {}, page_spec.get("right") or {}]
    if not any(halves):
        halves = [
            {"title": page_spec.get("left_title"), "bullets": page_spec.get("left_items")},
            {"title": page_spec.get("right_title"), "bullets": page_spec.get("right_items")},
        ]
    left, right = halves
    rect(slide, 0.58, 2.28, 7.15, 3.98, fill=THEME.paper)
    rect(slide, 7.98, 2.28, 4.77, 3.98, fill=THEME.navy)
    rect(slide, 0.58, 2.28, 0.09, 3.98, fill=THEME.accent)
    left_body = "\n".join(part for part in (_text(left.get("body")), _bullets(left.get("bullets"))) if part)
    right_body = "\n".join(part for part in (_text(right.get("body")), _bullets(right.get("bullets"))) if part)
    text_box(slide, 0.98, 2.75, 6.2, 0.62, _text(left.get("title")), size=25, bold=True, color=THEME.navy)
    text_box(slide, 0.98, 3.58, 6.15, 2.18, left_body, size=max(17, _body_size(left_body)), color=THEME.muted)
    text_box(slide, 8.38, 2.75, 3.92, 0.62, _text(right.get("title")), size=24, bold=True, color=THEME.accent)
    text_box(slide, 8.36, 3.38, 0.62, 0.68, "“", size=48, bold=True, color=THEME.accent, wrap=False)
    text_box(
        slide, 8.82, 3.72, 3.36, 1.72, right_body,
        size=24 if len(right_body) <= 48 else 21,
        bold=True, color=THEME.white, anchor=MSO_ANCHOR.MIDDLE,
    )


def render_process(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    items = _items(page_spec.get("items"))[:5]
    if not 3 <= len(items) <= 5:
        raise ValueError(f"slide {page} process requires 3 to 5 items")
    # Keep the first/last node centers inside their label columns.  When the
    # rail previously ran almost edge-to-edge, the end circles drifted away
    # from the centered titles and made the process look visually broken.
    left, right, axis_y = 1.75, 11.58, 3.45
    step = (right - left) / max(1, len(items) - 1)
    centers = [left + index * step for index in range(len(items))]
    # Draw directional edges first so arrowheads remain visible behind nodes
    # without adding four extra text boxes that inflate page density.
    for start, end in zip(centers, centers[1:]):
        line(slide, start + 0.31, axis_y, end - 0.31, axis_y, color=THEME.navy, weight=1.5, arrow=True)
    label_width = min(2.45, max(1.95, (right - left) / len(items) * 0.96))
    for index, (item, center) in enumerate(zip(items, centers)):
        circle(slide, center - 0.27, axis_y - 0.27, 0.54, fill=THEME.accent if index == 0 or index == len(items) - 1 else THEME.navy, line=THEME.white, line_width=1)
        text_box(slide, center - 0.27, axis_y - 0.17, 0.54, 0.24, f"{index + 1}", size=12, bold=True, color=THEME.white, align=PP_ALIGN.CENTER, wrap=False)
        body = _text(item.get("body")) or _text(item.get("desc")) or _text(item.get("detail"))
        label_x = max(0.58, min(center - label_width / 2, 12.75 - label_width))
        item_title = _text(item.get("title"))
        item_title_size = 18 if len(item_title) <= 9 else 15
        text_box(slide, label_x, 4.02, label_width, 0.58, item_title, size=item_title_size, bold=True, color=THEME.navy, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        text_box(slide, label_x, 4.78, label_width, 0.88, body, size=16, color=THEME.muted, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    laterals = [_text(value) for value in _items(page_spec.get("lateral")) if _text(value)]
    if laterals:
        band_y, band_h, band_x, band_w = 6.03, 0.62, left, right - left
        rect(slide, band_x, band_y, band_w, band_h, fill=THEME.paper, line=THEME.line, line_width=0.75)
        text_box(slide, band_x + 0.18, band_y + 0.08, 2.4, band_h - 0.16, "横向能力 →", size=14, bold=True, color=THEME.accent, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        segment_width = (band_w - 2.7) / max(1, len(laterals))
        for index, name in enumerate(laterals):
            x = band_x + 2.7 + index * segment_width
            circle(slide, x + 0.16, band_y + band_h / 2 - 0.16, 0.32, fill=THEME.navy, line=None)
            text_box(slide, x + 0.16, band_y + band_h / 2 - 0.17, 0.32, 0.32, str(index + 1), size=14, bold=True, color=THEME.white, align=PP_ALIGN.CENTER, wrap=False)
            text_box(slide, x + 0.54, band_y + 0.08, segment_width - 0.6, band_h - 0.16, name, size=15, bold=True, color=THEME.navy, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)


def render_metrics(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    items = (_items(page_spec.get("items")) or _items(page_spec.get("metrics")))[:4]
    if not items:
        raise ValueError(f"slide {page} metrics requires items")
    width = 11.9 / len(items)
    for index, item in enumerate(items):
        x = 0.72 + index * width
        value, value_size = _metric_value_style(_text(item.get("value")), width)
        text_box(slide, x, 2.42, width - 0.2, 1.12, value, size=value_size, bold=True, color=THEME.accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        line(slide, x + width * 0.25, 3.76, x + width * 0.75, 3.76, color=THEME.line, weight=1)
        text_box(slide, x, 4.02, width - 0.2, 0.7, _text(item.get("label")), size=18, bold=True, color=THEME.navy, align=PP_ALIGN.CENTER)
        body = _text(item.get("body")) or _text(item.get("note")) or _text(item.get("desc"))
        text_box(slide, x + 0.08, 4.9, width - 0.36, 1.0, body, size=_body_size(body, roomy=17, compact=16), color=THEME.muted, align=PP_ALIGN.CENTER)
        if index < len(items) - 1:
            line(slide, x + width - 0.1, 2.55, x + width - 0.1, 5.92, color=THEME.line, weight=1)


def render_chart(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    # Source provenance already has a dedicated low-key footer. Never promote
    # an internal path into the audience-facing insight card when the author did
    # not supply a real chart interpretation.
    note = _text(page_spec.get("note")) or _text(page_spec.get("intro"))
    chart_width = 8.35 if note else 12.0
    native_chart(slide, 0.68, 2.2, chart_width, 4.18, _items(page_spec.get("categories")), _items(page_spec.get("series")), chart_type=_text(page_spec.get("chart_type")) or "column", theme=THEME, show_legend=True, show_data_labels=True, label_size=13)
    if not note:
        return
    rect(slide, 9.3, 2.42, 3.38, 3.7, fill=THEME.navy)
    rect(slide, 9.3, 2.42, 0.08, 3.7, fill=THEME.accent)
    text_box(slide, 9.68, 2.82, 2.62, 0.55, _text(page_spec.get("note_title")) or "关键解读", size=20, bold=True, color=THEME.accent)
    text_box(slide, 9.68, 3.72, 2.62, 1.88, note, size=_body_size(note, roomy=17, compact=16), color=THEME.white)


def render_table(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    rows = _items(page_spec.get("rows"))[:6]
    editable_table(slide, 0.65, 2.3, 12.05, 4.0, _items(page_spec.get("headers")), rows, theme=THEME, header_size=16, body_size=16, cell_margin=0.055)


def render_matrix(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    points = _items(page_spec.get("points"))
    if points:
        matrix(slide, 0.9, 2.25, 11.55, 4.05, points, x_label=_text(page_spec.get("x_label")), y_label=_text(page_spec.get("y_label")), theme=THEME, label_size=14, axis_label_size=13, point_diameter=0.18, point_label_width=1.45)
        return
    # In business writing, “matrix” often means a comparison grid rather than
    # an x/y scatter. Accept that common compact schema without making the
    # content agent learn renderer internals.
    columns = _items(page_spec.get("columns")) or _items(page_spec.get("cols"))
    row_labels = _items(page_spec.get("rows"))
    values = _items(page_spec.get("values"))
    if columns and row_labels and len(row_labels) == len(values):
        rows = [[label, *list(value)] for label, value in zip(row_labels, values)]
        editable_table(slide, 0.55, 2.25, 12.2, 4.05, ["对象", *columns], rows, theme=THEME, header_size=16, body_size=16, cell_margin=0.045)
        return
    raise ValueError(f"slide {page} matrix requires points or columns/rows/values")


def render_timeline(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = _section_slide(deck, spec, page_spec, page, total)
    raw_items = _items(page_spec.get("items")) or _items(page_spec.get("milestones"))
    items = [_timeline_milestone(value, index) for index, value in enumerate(raw_items, 1)]
    if not 2 <= len(items) <= 6:
        raise ValueError(f"slide {page} timeline requires 2 to 6 milestones")
    timeline_diagram(
        slide, 0.68, 2.35, 12.0, 3.95, items, theme=THEME,
        label_size=16, title_size=18, detail_size=16, dot_diameter=0.2,
    )


def render_closing(deck, spec: dict, page_spec: dict, page: int, total: int) -> None:
    slide = new_slide(deck, THEME.navy)
    rect(slide, 0, 0, 13.333, 0.07, fill=THEME.accent)
    text_box(slide, 0.72, 0.62, 5.0, 0.34, _text(page_spec.get("kicker")) or "DECISION / NEXT STEP", size=16, bold=True, color=THEME.accent, wrap=False)
    title = _text(page_spec.get("title"))
    closing_size = _headline_size(title, 11.6)
    commitments = _closing_commitments(page_spec)
    if commitments:
        text_box(slide, 0.72, 1.15, 11.6, 0.78, title, size=min(closing_size, 36), bold=True, color=THEME.white)
        text_box(slide, 0.72, 2.00, 11.7, 0.54, _text(page_spec.get("takeaway")), size=18, color=RGBColor(0xD7, 0xE2, 0xF0))
        text_box(slide, 0.72, 2.68, 5.8, 0.3, _text(page_spec.get("decisions_label")) or "待确认决策", size=14, bold=True, color=THEME.accent, wrap=False)
        items = _closing_actions(page_spec)
        for index, item in enumerate(items):
            x = 0.72 + index * 4.05
            if index:
                line(slide, x - 0.2, 3.08, x - 0.2, 4.36, color=RGBColor(0x46, 0x5A, 0x70), weight=1)
            text_box(slide, x, 3.06, 0.45, 0.28, f"0{index + 1}", size=14, bold=True, color=THEME.accent, wrap=False)
            text_box(slide, x, 3.48, 3.55, 0.36, _text(item.get("title")), size=17, bold=True, color=THEME.white)
            body = _text(item.get("body"))
            text_box(slide, x, 3.93, 3.55, 0.48, body, size=16, color=RGBColor(0xD7, 0xE2, 0xF0))
        line(slide, 0.72, 4.60, 12.58, 4.60, color=RGBColor(0x46, 0x5A, 0x70), weight=1)
        text_box(slide, 0.72, 4.78, 6.5, 0.3, _text(page_spec.get("commitments_label")) or "阶段结果承诺", size=14, bold=True, color=THEME.accent, wrap=False)
        for index, item in enumerate(commitments):
            x = 0.72 + index * 4.05
            if index:
                line(slide, x - 0.2, 5.20, x - 0.2, 6.48, color=RGBColor(0x46, 0x5A, 0x70), weight=1)
            text_box(slide, x, 5.20, 3.55, 0.38, _text(item.get("title")), size=20, bold=True, color=THEME.accent)
            body = _text(item.get("body"))
            text_box(slide, x, 5.72, 3.55, 0.72, body, size=16, color=THEME.white)
        note = _text(page_spec.get("note"))
        if note:
            text_box(slide, 0.72, 6.72, 11.85, 0.34, note, size=16, color=RGBColor(0xD7, 0xE2, 0xF0), wrap=False)
        return
    text_box(slide, 0.72, 1.45, 11.6, 1.35, title, size=closing_size, bold=True, color=THEME.white)
    text_box(slide, 0.72, 2.98, 11.0, 0.72, _text(page_spec.get("takeaway")), size=20, color=RGBColor(0xD7, 0xE2, 0xF0))
    items = _closing_actions(page_spec)
    for index, item in enumerate(items):
        x = 0.72 + index * 4.05
        if index:
            line(slide, x - 0.2, 4.35, x - 0.2, 6.2, color=RGBColor(0x46, 0x5A, 0x70), weight=1)
        text_box(slide, x, 4.32, 0.45, 0.3, f"0{index + 1}", size=14, bold=True, color=THEME.accent, wrap=False)
        text_box(slide, x, 4.92, 3.55, 0.42, _text(item.get("title")), size=18, bold=True, color=THEME.accent)
        body = _text(item.get("body"))
        text_box(slide, x, 5.52, 3.55, 0.72, body, size=_body_size(body, roomy=16, compact=16), color=THEME.white)


RENDERERS = {
    "cover": render_cover,
    "statement": render_statement,
    "cards": render_cards,
    "split": render_split,
    "process": render_process,
    "metrics": render_metrics,
    "chart": render_chart,
    "table": render_table,
    "matrix": render_matrix,
    "timeline": render_timeline,
    "closing": render_closing,
}


# === TASK PAGE COMPOSITIONS: BEGIN ===
# Scaffolded formal decks must replace this empty registry with one named,
# task-local renderer for every slide number.  Each renderer has the uniform
# signature (deck, spec, page_spec, page, total).  It may reuse the stable
# mechanics above, but it should choose geometry, hierarchy, and focal scale
# for that page's own narrative job rather than dispatching the whole deck by
# ``kind``.  The compact build_task_pptx.py route substitutes stable mode and
# intentionally keeps this registry empty.
TASK_PAGE_RENDERERS = {}
# === TASK PAGE COMPOSITIONS: END ===


def _validate_task_page_renderers(total: int) -> None:
    if COMPOSITION_MODE != "per-slide":
        return
    expected = set(range(1, total + 1))
    actual = set(TASK_PAGE_RENDERERS)
    if actual != expected:
        missing = sorted(expected - actual)
        extra = sorted(actual - expected)
        raise ValueError(
            "per-slide customization is incomplete; "
            f"missing renderer pages={missing}, unexpected pages={extra}"
        )
    invalid = sorted(page for page, renderer in TASK_PAGE_RENDERERS.items() if not callable(renderer))
    if invalid:
        raise ValueError(f"per-slide renderers must be callable for pages={invalid}")
    generic_renderers = set(RENDERERS.values())
    thin_wrappers = []
    for page, renderer in TASK_PAGE_RENDERERS.items():
        code = getattr(renderer, "__code__", None)
        names = set(getattr(code, "co_names", ()))
        if renderer in generic_renderers or names.intersection({"RENDERERS", "_render_page"}):
            thin_wrappers.append(page)
    if thin_wrappers:
        raise ValueError(
            "per-slide customization cannot be thin dispatcher wrappers around "
            "the generic RENDERERS registry; implement page-owned geometry for "
            f"pages={sorted(thin_wrappers)}"
        )


def _validate_visual_realization(deck) -> None:
    """Ensure the declared primary proof object became a real slide object."""

    if COMPOSITION_MODE != "per-slide":
        return
    for page, slide in enumerate(deck.slides, 1):
        visual = SLIDE_VISUALS.get(page, {})
        asset = visual.get("asset") if isinstance(visual, dict) else None
        asset_type = _text(asset.get("type")).lower() if isinstance(asset, dict) else ""
        if asset_type in {"photo", "illustration", "screenshot"}:
            if not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes):
                raise ValueError(f"slide {page} did not realize its planned media asset as a picture")
        elif asset_type == "chart":
            if not any(getattr(shape, "has_chart", False) for shape in slide.shapes):
                raise ValueError(f"slide {page} did not realize its planned native chart")
        elif asset_type == "table":
            if not any(getattr(shape, "has_table", False) for shape in slide.shapes):
                raise ValueError(f"slide {page} did not realize its planned native table")


def _apply_planned_typeface(deck) -> None:
    """Apply the selected font pair to all editable text after composition."""

    if not VISUAL_DIRECTION:
        return
    typography = VISUAL_DIRECTION.get("typography")
    if not isinstance(typography, dict):
        return
    font_cn = _text(typography.get("font_cn")) or THEME.font_cn
    font_en = _text(typography.get("font_en")) or THEME.font_en
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = font_en
                    properties = run._r.get_or_add_rPr()
                    east_asia = properties.find(qn("a:ea"))
                    if east_asia is None:
                        east_asia = etree.SubElement(properties, qn("a:ea"))
                    east_asia.set("typeface", font_cn)
def build() -> Path:
    global THEME, VISUAL_PLAN, VISUAL_DIRECTION, SLIDE_VISUALS
    spec = load_spec()
    slides = _items(spec.get("slides"))
    total = len(slides)
    VISUAL_PLAN = _load_visual_plan(total)
    VISUAL_DIRECTION = VISUAL_PLAN.get("selected_direction", {}) if VISUAL_PLAN else {}
    SLIDE_VISUALS = {
        int(item["slide"]): item
        for item in _items(VISUAL_PLAN.get("slides"))
        if isinstance(item, dict) and isinstance(item.get("slide"), int)
    }
    THEME = _theme_for_spec(spec, VISUAL_DIRECTION)
    deck = new_presentation()
    _validate_task_page_renderers(total)
    for page, page_spec in enumerate(slides, 1):
        if COMPOSITION_MODE == "per-slide":
            TASK_PAGE_RENDERERS[page](deck, spec, page_spec, page, total)
        else:
            kind = _text(page_spec.get("kind")).lower()
            renderer = RENDERERS[kind]
            if kind == "cover":
                renderer(deck, spec, page_spec)
            else:
                renderer(deck, spec, page_spec, page, total)
    _validate_visual_realization(deck)
    _apply_planned_typeface(deck)
    assert_deck_bounds(deck)
    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(OUTPUT_PPTX))
    return OUTPUT_PPTX


if __name__ == "__main__":
    print(build())
