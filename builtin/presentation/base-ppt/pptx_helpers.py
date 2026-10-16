#!/usr/bin/env python3
"""Small, path-independent python-pptx helpers for bespoke OpenKitty decks.

This module is intentionally content-agnostic.  It owns the mechanics that are
easy to get wrong repeatedly (canvas setup, colors, text fitting, image aspect
ratio, editable tables/charts, common semantic diagrams, and bounds checks),
while the task script remains free to design each slide for its actual story.

Copy this file next to a task-specific ``build_ppt.py`` and import it locally;
no OpenKitty installation path is required at execution time.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Sequence

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


WIDE_WIDTH = 13.333
WIDE_HEIGHT = 7.5


@dataclass(frozen=True)
class Theme:
    name: str
    background: str
    surface: str
    dark: str
    primary: str
    secondary: str
    text: str
    muted: str
    border: str
    on_dark: str
    series: tuple[str, ...]
    heading_font: str = "Microsoft YaHei"
    body_font: str = "Microsoft YaHei"


THEMES = {
    "navy_gold": Theme(
        "navy_gold", "FFFFFF", "F5F3EE", "051C2C", "BFA960", "1B3A5C",
        "051C2C", "6F6F6F", "DEDCD6", "FFFFFF",
        ("051C2C", "C0392B", "4E8C4D", "8064A2", "4CA6BF", "BFA960"),
    ),
    "google_clean": Theme(
        "google_clean", "FFFFFF", "F8F9FA", "202124", "4285F4", "34A853",
        "202124", "5F6368", "DADCE0", "FFFFFF",
        ("4285F4", "EA4335", "FBBC04", "34A853"),
        heading_font="Arial", body_font="Arial",
    ),
    "warm_editorial": Theme(
        "warm_editorial", "F7F0E8", "FFF8F0", "2B2118", "D97757", "C9A227",
        "2B2118", "7C6F64", "E4D7C8", "F7F0E8",
        ("D97757", "5E7D69", "C9A227", "7C5C6E"),
        heading_font="Georgia", body_font="Calibri",
    ),
}

THEME_ALIASES = {
    "google_style": "google_clean",
    "anthropic": "warm_editorial",
}


def rgb(value: str | RGBColor) -> RGBColor:
    if isinstance(value, RGBColor):
        return value
    value = str(value).strip().lstrip("#")[:6]
    if len(value) != 6:
        raise ValueError(f"expected six-digit hex color, got {value!r}")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def resolve_theme(theme: str | Theme) -> Theme:
    if isinstance(theme, Theme):
        return theme
    name = THEME_ALIASES.get(str(theme).strip().lower(), str(theme).strip().lower())
    if name not in THEMES:
        raise ValueError(f"unknown helper theme: {theme}")
    return THEMES[name]


def visual_units(text: object) -> float:
    units = 0.0
    for char in str(text or ""):
        code = ord(char)
        if char.isspace():
            units += 0.30
        elif 0x2E80 <= code <= 0x9FFF or 0xF900 <= code <= 0xFAFF:
            units += 1.0
        elif char.isupper():
            units += 0.68
        else:
            units += 0.54
    return units


def fitted_font_size(text: object, width: float, maximum: float, minimum: float, max_lines: int = 1) -> float:
    units = max(1.0, visual_units(text))
    capacity = max(1.0, width * 72.0 * max_lines / (maximum * 0.72))
    if units <= capacity:
        return maximum
    return round(max(minimum, maximum * capacity / units), 1)


class Deck:
    """A thin drawing surface for per-deck scripts, not an auto-layout engine."""

    def __init__(self, theme: str | Theme = "navy_gold", width: float = WIDE_WIDTH, height: float = WIDE_HEIGHT):
        self.theme = resolve_theme(theme)
        self.width = float(width)
        self.height = float(height)
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.width)
        self.prs.slide_height = Inches(self.height)

    def add_slide(self, background: str | None = None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.set_background(slide, background or self.theme.background)
        return slide

    @staticmethod
    def set_background(slide, color: str):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(color)

    @staticmethod
    def rect(slide, x: float, y: float, w: float, h: float, fill: str, line: str | None = None, line_width: float = 0.6):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_width)
        else:
            shape.line.fill.background()
        return shape

    def text(
        self,
        slide,
        x: float,
        y: float,
        w: float,
        h: float,
        text: object,
        size: float = 14,
        *,
        color: str | None = None,
        bold: bool = False,
        italic: bool = False,
        font: str | None = None,
        align=PP_ALIGN.LEFT,
        vertical=MSO_ANCHOR.TOP,
        margin: float = 4,
        fit: tuple[float, float, int] | None = None,
        wrap: bool = True,
    ):
        if fit:
            size = fitted_font_size(text, w, fit[0], fit[1], fit[2])
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = wrap
        frame.vertical_anchor = vertical
        frame.margin_left = frame.margin_right = Pt(margin)
        frame.margin_top = frame.margin_bottom = Pt(max(1, margin / 2))
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = str(text or "")
        run.font.name = font or self.theme.body_font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color or self.theme.text)
        return box

    def image_contain(
        self,
        slide,
        path: str | Path,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        background: str | None = None,
        border: str | None = None,
        caption: str = "",
    ):
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        caption_h = 0.30 if caption else 0.0
        image_h = max(0.20, h - caption_h)
        self.rect(slide, x, y, w, image_h, background or self.theme.background, border or self.theme.border)
        with Image.open(path) as image:
            pixel_w, pixel_h = image.size
        if pixel_w <= 0 or pixel_h <= 0:
            raise ValueError(f"invalid image dimensions: {path}")
        image_ratio = pixel_w / pixel_h
        box_ratio = w / image_h
        if image_ratio >= box_ratio:
            draw_w, draw_h = w, w / image_ratio
        else:
            draw_h, draw_w = image_h, image_h * image_ratio
        draw_x = x + (w - draw_w) / 2
        draw_y = y + (image_h - draw_h) / 2
        picture = slide.shapes.add_picture(str(path), Inches(draw_x), Inches(draw_y), Inches(draw_w), Inches(draw_h))
        if caption:
            self.text(slide, x, y + image_h + 0.02, w, caption_h - 0.02, caption, 9,
                      color=self.theme.muted, align=PP_ALIGN.CENTER, wrap=False)
        return picture

    def header(self, slide, title: str, *, eyebrow: str = "", page: int | None = None, insight: str = ""):
        t = self.theme
        if t.name == "google_clean":
            colors = t.series
            starts = (0.0, 0.25, 0.45, 0.60)
            widths = (0.25, 0.20, 0.15, 0.40)
            for start, width, color in zip(starts, widths, colors):
                self.rect(slide, self.width * start, 0, self.width * width, 0.05, color)
            title_color = t.text
        else:
            self.rect(slide, 0, 0, self.width, 1.02, t.dark)
            self.rect(slide, 0, 1.02, self.width, 0.035, t.primary)
            title_color = t.on_dark
        if eyebrow:
            self.text(slide, 0.62, 0.16, 8.4, 0.18, eyebrow.upper(), 9, color=t.primary, bold=True, wrap=False)
        title_y = 0.37 if eyebrow else 0.25
        title_w = 8.65 if insight else 11.7
        self.text(slide, 0.62, title_y, title_w, 0.48, title, 24, color=title_color,
                  bold=True, font=t.heading_font, fit=(24, 18, 2))
        if insight:
            rail_x = 9.55
            self.rect(slide, rail_x, 0.20, 0.035, 0.58, t.primary)
            self.text(slide, rail_x + 0.14, 0.20, 2.28, 0.58, insight, 9.5,
                      color=t.muted if t.name == "google_clean" else "D9DEE2")
        if page is not None:
            self.text(slide, 12.15, 0.13, 0.55, 0.20, f"{page:02d}", 8,
                      color=t.primary, bold=True, align=PP_ALIGN.RIGHT, wrap=False)

    def footer(self, slide, label: str = "", page: int | None = None):
        t = self.theme
        y = self.height - 0.30
        fill = t.dark if t.name == "navy_gold" else t.surface
        self.rect(slide, 0, y, self.width, 0.30, fill)
        self.rect(slide, 0, y, self.width, 0.012, t.primary if t.name == "navy_gold" else t.border)
        if label:
            self.text(slide, 0.62, y + 0.06, 8.5, 0.17, label, 8,
                      color="D9DEE2" if t.name == "navy_gold" else t.muted, wrap=False)
        if page is not None:
            self.text(slide, 11.9, y + 0.06, 0.80, 0.17, f"{page:02d}", 8,
                      color=t.primary, bold=True, align=PP_ALIGN.RIGHT, wrap=False)

    def cover(self, title: str, subtitle: str = "", *, eyebrow: str = "", tags: Sequence[str] = ()):
        t = self.theme
        background = t.dark if t.name != "google_clean" else t.background
        slide = self.add_slide(background)
        self.rect(slide, 0, 0, self.width, 0.05, t.primary)
        self.rect(slide, 0.68, 1.25, 0.06, 3.05, t.primary)
        title_color = t.on_dark if t.name != "google_clean" else t.text
        muted_color = "D9DEE2" if t.name != "google_clean" else t.muted
        if eyebrow:
            self.text(slide, 0.93, 1.25, 7.7, 0.22, eyebrow.upper(), 10,
                      color=t.primary, bold=True, wrap=False)
        self.text(slide, 0.93, 1.78, 8.0, 1.55, title, 36, color=title_color,
                  bold=True, font=t.heading_font, fit=(36, 25, 2))
        if subtitle:
            self.text(slide, 0.93, 3.43, 7.6, 0.64, subtitle, 15, color=muted_color)
        tag_x = 0.93
        for tag in tags[:4]:
            width = min(1.8, max(0.75, visual_units(tag) * 0.11 + 0.28))
            self.rect(slide, tag_x, 4.55, min(0.62, width * 0.35), 0.03, t.primary)
            self.text(slide, tag_x, 4.68, width, 0.22, tag, 9, color=muted_color, wrap=False)
            tag_x += width + 0.22
        return slide

    def kpi_strip(self, slide, metrics: Sequence[Mapping[str, object]], x: float, y: float, w: float, h: float):
        if not metrics:
            raise ValueError("kpi_strip requires metrics")
        gap = 0.14
        card_w = (w - gap * (len(metrics) - 1)) / len(metrics)
        for index, metric in enumerate(metrics):
            card_x = x + index * (card_w + gap)
            color = self.theme.series[index % len(self.theme.series)]
            fill = self.theme.dark if self.theme.name == "navy_gold" else self.theme.background
            self.rect(slide, card_x, y, card_w, h, fill, self.theme.border)
            self.rect(slide, card_x, y, card_w, 0.035, color)
            value = str(metric.get("value", ""))
            label = str(metric.get("label", ""))
            detail = str(metric.get("detail", metric.get("sub", "")))
            value_color = self.theme.on_dark if self.theme.name == "navy_gold" else color
            label_color = self.theme.primary if self.theme.name == "navy_gold" else self.theme.text
            self.text(slide, card_x + 0.10, y + 0.18, card_w - 0.20, h * 0.38, value, 30,
                      color=value_color, bold=True, font=self.theme.heading_font,
                      align=PP_ALIGN.CENTER, fit=(30, 18, 1), wrap=False)
            self.text(slide, card_x + 0.10, y + h * 0.55, card_w - 0.20, h * 0.18, label, 11,
                      color=label_color, bold=True, align=PP_ALIGN.CENTER)
            if detail:
                self.text(slide, card_x + 0.10, y + h * 0.78, card_w - 0.20, h * 0.14, detail, 8.5,
                          color="D9DEE2" if self.theme.name == "navy_gold" else self.theme.muted,
                          align=PP_ALIGN.CENTER)

    def timeline(self, slide, items: Sequence[Mapping[str, object]], x: float, y: float, w: float, h: float):
        if len(items) < 2:
            raise ValueError("timeline requires at least two milestones")
        axis_y = y + h * 0.50
        left, right = x + 0.28, x + w - 0.28
        self.rect(slide, left, axis_y, right - left, 0.025, self.theme.primary)
        column_w = (right - left) / len(items)
        for index, item in enumerate(items):
            center_x = left + column_w * (index + 0.5)
            color = self.theme.series[index % len(self.theme.series)]
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center_x - 0.08), Inches(axis_y - 0.07), Inches(0.16), Inches(0.16))
            dot.fill.solid(); dot.fill.fore_color.rgb = rgb(color)
            dot.line.color.rgb = rgb(self.theme.background); dot.line.width = Pt(1.0)
            top = index % 2 == 0
            block_y = y + 0.06 if top else axis_y + 0.28
            self.text(slide, center_x - column_w * 0.45, block_y, column_w * 0.90, 0.24,
                      item.get("label", item.get("year", "")), 10, color=self.theme.primary,
                      bold=True, align=PP_ALIGN.CENTER, wrap=False)
            self.text(slide, center_x - column_w * 0.47, block_y + 0.28, column_w * 0.94, 0.34,
                      item.get("title", ""), 11, color=self.theme.text, bold=True,
                      align=PP_ALIGN.CENTER, font=self.theme.heading_font)
            self.text(slide, center_x - column_w * 0.47, block_y + 0.66, column_w * 0.94, 0.48,
                      item.get("detail", ""), 8.5, color=self.theme.muted, align=PP_ALIGN.CENTER)

    def matrix(
        self,
        slide,
        items: Sequence[Mapping[str, object]],
        quadrants: Sequence[str],
        x_label: str,
        y_label: str,
        x: float,
        y: float,
        w: float,
        h: float,
    ):
        if len(quadrants) != 4:
            raise ValueError("matrix requires four quadrant labels")
        plot_x, plot_y = x + 0.55, y + 0.12
        plot_w, plot_h = w - 0.75, h - 0.52
        half_w, half_h = plot_w / 2, plot_h / 2
        positions = ((plot_x, plot_y + half_h), (plot_x + half_w, plot_y + half_h),
                     (plot_x, plot_y), (plot_x + half_w, plot_y))
        for index, (qx, qy) in enumerate(positions):
            fill = self.theme.surface if index in (0, 3) else self.theme.background
            self.rect(slide, qx, qy, half_w, half_h, fill, self.theme.border)
            self.text(slide, qx + 0.10, qy + 0.08, half_w - 0.20, 0.24, quadrants[index], 10,
                      color=self.theme.text if index in (0, 3) else self.theme.muted, bold=True)
        self.rect(slide, plot_x + half_w - 0.01, plot_y, 0.02, plot_h, self.theme.primary)
        self.rect(slide, plot_x, plot_y + half_h - 0.01, plot_w, 0.02, self.theme.primary)
        self.text(slide, plot_x, plot_y + plot_h + 0.08, plot_w, 0.24, x_label, 9,
                  color=self.theme.muted, bold=True, align=PP_ALIGN.CENTER, wrap=False)
        self.text(slide, x, plot_y, 0.50, 0.40, y_label, 9, color=self.theme.muted,
                  bold=True, align=PP_ALIGN.CENTER)
        for index, item in enumerate(items):
            point_x = float(item["x"]); point_y = float(item["y"])
            if not 0 <= point_x <= 1 or not 0 <= point_y <= 1:
                raise ValueError("matrix item x/y must be between 0 and 1")
            center_x = plot_x + plot_w * point_x
            center_y = plot_y + plot_h * (1 - point_y)
            color = self.theme.series[index % len(self.theme.series)]
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center_x - 0.07), Inches(center_y - 0.07), Inches(0.14), Inches(0.14))
            dot.fill.solid(); dot.fill.fore_color.rgb = rgb(color)
            dot.line.color.rgb = rgb(self.theme.background); dot.line.width = Pt(0.8)
            self.text(slide, min(plot_x + plot_w - 1.32, center_x + 0.09), center_y - 0.12, 1.25, 0.24,
                      item.get("label", ""), 8.5, color=self.theme.text, bold=True, wrap=False)

    def table(self, slide, headers: Sequence[object], rows: Sequence[Sequence[object]], x: float, y: float, w: float, h: float):
        if not headers or not rows:
            raise ValueError("table requires headers and rows")
        row_count = len(rows) + 1
        col_count = len(headers)
        table = slide.shapes.add_table(row_count, col_count, Inches(x), Inches(y), Inches(w), Inches(h)).table
        for column, header in enumerate(headers):
            cell = table.cell(0, column)
            cell.text = str(header)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(self.theme.dark)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.name = self.theme.heading_font
                paragraph.font.size = Pt(10)
                paragraph.font.bold = True
                paragraph.font.color.rgb = rgb(self.theme.on_dark)
        for row_index, row in enumerate(rows, 1):
            for column in range(col_count):
                cell = table.cell(row_index, column)
                cell.text = str(row[column]) if column < len(row) else ""
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(self.theme.surface if row_index % 2 else self.theme.background)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = self.theme.body_font
                    paragraph.font.size = Pt(9)
                    paragraph.font.color.rgb = rgb(self.theme.text)
                    paragraph.alignment = PP_ALIGN.LEFT if column else PP_ALIGN.CENTER
                cell.margin_left = cell.margin_right = Pt(4)
                cell.margin_top = cell.margin_bottom = Pt(2)
        return table

    def native_chart(
        self,
        slide,
        categories: Sequence[object],
        series: Sequence[Mapping[str, object]],
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        chart_type: str = "column",
        title: str = "",
    ):
        if not categories or not series:
            raise ValueError("native_chart requires categories and series")
        chart_types = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "bar_100": XL_CHART_TYPE.BAR_STACKED_100,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }
        if chart_type not in chart_types:
            raise ValueError(f"unsupported native chart type: {chart_type}")
        data = CategoryChartData()
        data.categories = [str(value) for value in categories]
        for item in series:
            data.add_series(str(item.get("name", "")), list(item.get("values", [])))
        chart = slide.shapes.add_chart(chart_types[chart_type], Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        chart.has_title = bool(title)
        if title:
            chart.chart_title.text_frame.text = title
            for paragraph in chart.chart_title.text_frame.paragraphs:
                paragraph.font.name = self.theme.heading_font
                paragraph.font.size = Pt(13)
                paragraph.font.bold = True
                paragraph.font.color.rgb = rgb(self.theme.text)
        chart.has_legend = len(series) > 1 or chart_type in {"pie", "doughnut"}
        for index, chart_series in enumerate(chart.series):
            color = self.theme.series[index % len(self.theme.series)]
            try:
                chart_series.format.fill.solid(); chart_series.format.fill.fore_color.rgb = rgb(color)
                chart_series.format.line.color.rgb = rgb(color)
            except (AttributeError, TypeError, ValueError):
                pass
        if chart_type in {"pie", "doughnut"} and chart.series:
            for index, point in enumerate(chart.series[0].points):
                color = self.theme.series[index % len(self.theme.series)]
                point.format.fill.solid(); point.format.fill.fore_color.rgb = rgb(color)
                point.format.line.color.rgb = rgb(self.theme.background)
        try:
            chart.plots[0].has_data_labels = True
            chart.plots[0].data_labels.font.size = Pt(8)
            chart.plots[0].data_labels.font.color.rgb = rgb(self.theme.text)
        except (AttributeError, TypeError, ValueError):
            pass
        return chart

    def validate(self) -> list[str]:
        issues = []
        slide_width = self.prs.slide_width
        slide_height = self.prs.slide_height
        for slide_index, slide in enumerate(self.prs.slides, 1):
            if not slide.shapes:
                issues.append(f"slide {slide_index}: empty slide")
            for shape in slide.shapes:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_width or shape.top + shape.height > slide_height:
                    issues.append(f"slide {slide_index}: shape outside canvas: {shape.name}")
        return issues

    def save(self, path: str | Path, *, strict_bounds: bool = True) -> Path:
        issues = self.validate()
        if strict_bounds and issues:
            raise ValueError("; ".join(issues))
        output = Path(path)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(output)
        return output


__all__ = [
    "Deck", "Theme", "THEMES", "THEME_ALIASES", "WIDE_WIDTH", "WIDE_HEIGHT",
    "fitted_font_size", "resolve_theme", "rgb", "visual_units",
]
