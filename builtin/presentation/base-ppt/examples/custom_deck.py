#!/usr/bin/env python3
"""Copyable bespoke-deck example built on base-ppt/pptx_helpers.py.

For a real task, copy this file and ``pptx_helpers.py`` into the task storage
directory, then replace the demo content and page compositions.  The import
works both inside the skill directory and when both files are copied together.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR


HERE = Path(__file__).resolve().parent
for candidate in (HERE, HERE.parent):
    if (candidate / "pptx_helpers.py").is_file():
        sys.path.insert(0, str(candidate))
        break
else:
    raise SystemExit("Copy pptx_helpers.py next to this script before running it.")

from pptx_helpers import Deck  # noqa: E402


def content_slide(deck: Deck, title: str, eyebrow: str, page: int, insight: str = ""):
    slide = deck.add_slide()
    deck.header(slide, title, eyebrow=eyebrow, page=page, insight=insight)
    deck.footer(slide, "Custom renderer example · replace with a real source label", page)
    return slide


def build(output: Path, theme: str):
    deck = Deck(theme)
    t = deck.theme

    deck.cover(
        "从证据到行动的项目叙事",
        "演示数据 · 展示如何用稳定 helper 快速编写每份专用 PPT 脚本",
        eyebrow="CUSTOM PYTHON-PPTX REFERENCE",
        tags=["证据", "判断", "行动"],
    )

    slide = content_slide(deck, "增长数据只有与解释放在一起才形成判断", "01 · EVIDENCE", 2)
    deck.kpi_strip(slide, [
        {"value": "18,010", "label": "基期样本", "detail": "演示值"},
        {"value": "+128%", "label": "阶段增速", "detail": "演示值"},
        {"value": "3", "label": "备选方向", "detail": "演示值"},
    ], 0.62, 1.35, 12.08, 1.32)
    deck.native_chart(
        slide,
        ["2024", "2026", "2028", "2030"],
        [{"name": "规模", "values": [18, 46, 92, 155]}],
        0.62, 2.98, 8.05, 3.72,
        chart_type="column",
        title="示例规模走势（演示数据）",
    )
    deck.rect(slide, 9.02, 2.98, 3.68, 3.72, t.surface, t.border)
    deck.rect(slide, 9.02, 2.98, 0.05, 3.72, t.primary)
    deck.text(slide, 9.28, 3.24, 3.05, 0.36, "为什么重要", 17,
              color=t.text, bold=True, font=t.heading_font)
    deck.rect(slide, 9.28, 3.78, 0.72, 0.03, t.primary)
    deck.text(slide, 9.28, 4.03, 3.05, 1.78,
              "• 不只展示增长，还说明增长发生在哪个阶段。\n\n"
              "• 把图表、结论与下一步放在同一阅读路径。\n\n"
              "• 每个数字都带口径、时期或来源",
              12, color=t.text)

    slide = content_slide(deck, "四个里程碑把早期验证推进到稳定交付", "02 · ROADMAP", 3)
    deck.timeline(slide, [
        {"label": "2024", "title": "原型验证", "detail": "验证核心机制"},
        {"label": "2026", "title": "场景试点", "detail": "验证可靠性"},
        {"label": "2028", "title": "量产爬坡", "detail": "稳定供应链"},
        {"label": "2030", "title": "规模交付", "detail": "扩大场景覆盖"},
    ], 0.62, 1.45, 12.08, 5.25)

    slide = content_slide(deck, "价值与成熟度共同决定优先级", "03 · PRIORITY", 4,
                          "矩阵坐标必须来自真实评价或明确判断，不能由脚本猜测。")
    deck.matrix(
        slide,
        [
            {"label": "方案 A", "x": 0.82, "y": 0.74},
            {"label": "方案 B", "x": 0.58, "y": 0.52},
            {"label": "方案 C", "x": 0.28, "y": 0.86},
        ],
        ["持续观察", "选择性投入", "能力储备", "优先规模化"],
        "商业成熟度", "战略价值",
        0.62, 1.35, 12.08, 5.50,
    )

    slide = content_slide(deck, "精确比较交给表格，决策含义放在侧栏", "04 · COMPARISON", 5)
    deck.table(
        slide,
        ["方案", "成熟度", "部署成本", "建议"],
        [
            ["方案 A", "高", "中", "优先试点"],
            ["方案 B", "中", "低", "限定场景"],
            ["方案 C", "低", "高", "持续跟踪"],
        ],
        0.62, 1.55, 8.20, 3.05,
    )
    deck.rect(slide, 0.62, 4.94, 8.20, 1.23, t.surface, t.border)
    deck.rect(slide, 0.62, 4.94, 0.05, 1.23, t.primary)
    deck.text(slide, 0.88, 5.18, 7.55, 0.66,
              "表格保留精确值；页面标题与右侧结论负责告诉听众如何使用这些值。",
              13, color=t.text, bold=True, vertical=MSO_ANCHOR.MIDDLE)
    deck.rect(slide, 9.15, 1.42, 3.55, 4.75, t.surface, t.border)
    deck.rect(slide, 9.15, 1.42, 0.05, 4.75, t.primary)
    deck.text(slide, 9.42, 1.76, 2.95, 0.40, "决策结论", 18,
              color=t.text, bold=True, font=t.heading_font)
    deck.text(slide, 9.42, 2.46, 2.95, 2.72,
              "方案 A 的成熟度已经足以支撑试点；方案 B 保留在明确边界内；方案 C 暂不进入规模投入。",
              14, color=t.text)

    slide = deck.add_slide(t.dark)
    deck.rect(slide, 0, 0, deck.width, 0.05, t.primary)
    deck.rect(slide, 0.68, 1.42, 0.06, 2.10, t.primary)
    deck.text(slide, 0.96, 1.42, 8.50, 1.02, "把证据转化为下一步动作", 34,
              color=t.on_dark, bold=True, font=t.heading_font, fit=(34, 25, 2))
    deck.text(slide, 0.96, 2.72, 7.70, 0.64,
              "明确试点场景、验证口径和复盘时间，再进入下一轮投入决策。",
              15, color="D9DEE2")
    deck.rect(slide, 0, 6.34, deck.width, 1.16, t.secondary)
    deck.rect(slide, 0, 6.34, deck.width, 0.03, t.primary)

    issues = deck.validate()
    if issues:
        raise ValueError("helper validation failed: " + "; ".join(issues))
    saved = deck.save(output)
    print(f"custom deck saved: {saved} ({len(deck.prs.slides)} slides)")


def main():
    parser = argparse.ArgumentParser(description="Build a copyable bespoke PPT example.")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--theme", default="navy_gold")
    args = parser.parse_args()
    build(args.output, args.theme)


if __name__ == "__main__":
    main()
